"""Offline repair of patch16 tables; never invoke model inference or open test data."""
import copy
import csv
import html
import math
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.evaluate_pemsbay_patch16 import check, output, read, write, sha, MODELS, GROUPS, METRICS


def rank_values(values):
    assert all(math.isfinite(v) for v in values)
    unique = sorted(set(values))
    return [unique.index(v)+1 for v in values]


def publish(records, split, plan):
    out = ROOT / f'paper/{split}_seed32_hparam_patch16_20260918'
    out.mkdir(parents=True, exist_ok=True)
    notes = [f'{split}，METR-LA/PEMS-BAY，seed32，Focal＋九个基线。',
             '仅PEMS-BAY Focal更新为patch16；另外19项原split记录原样复用。原Focal表保留。',
             '候选按预定validation门槛选择；这是已查看原test之后的补充探索，后续test不是未受观察影响的独立确认。',
             '第一名加粗、第二个不同数值下划线；精确并列共享名次，按未舍入值排序。',
             'SLPF为单阶段预测适配；STBP为离线固定图适配；ST-FiT保留源Gumbel随机评估。',
             'MAE/RMSE为原速度单位，MAPE为%；不是多种子显著性结论。']
    wb = Workbook(); ws = wb.active; ws.title = '说明'
    for note in notes: ws.append([note])
    ws.column_dimensions['A'].width = 130
    md = [f'# {split} seed32：PEMS-BAY Focal patch16', '', *notes, '']
    sections, cells, checks, rankings = [], [], [], {}
    for ds, items in records.items():
        assert set(items) == set(MODELS)
        rankings[ds] = {}
        for group, conditions in GROUPS.items():
            title = ds + ('主表' if group == 'main' else '阶段表')
            keys = [(c,m) for c in conditions for m in METRICS]
            headers = ['模型'] + [f'{c}/{m}' for c,m in keys]
            sheet = wb.create_sheet(title); sheet.append(headers)
            ranks = {k: rank_values([items[n]['table'][group][k[0]][k[1]] for n in MODELS]) for k in keys}
            md += ['## '+title, '', '| '+' | '.join(headers)+' |', '|'+'|'.join(['---']+['---:']*len(keys))+'|']
            body = []; rankings[ds][group] = {}
            for ri, name in enumerate(MODELS, 2):
                label = ('ALFNet（Focal，patch16）' if ds=='PEMSBAY' else 'ALFNet（Focal）') if name=='ALFNet' else {'STFiT':'ST-FiT','SLPF':'SLPF（单阶段）'}.get(name,name)
                sheet.cell(ri,1,label); checks.append((title,ri,1,label,None))
                shown = []; places = []
                for ci,(condition,metric) in enumerate(keys,2):
                    value = items[name]['table'][group][condition][metric]
                    rank = ranks[(condition,metric)][ri-2]; places.append(rank)
                    cell = sheet.cell(ri,ci,value); cell.number_format='0.000'
                    cell.font=Font(bold=rank==1,underline='single' if rank==2 else None)
                    checks.append((title,ri,ci,value,rank))
                    cells.append([ds,name,group,condition,metric,value,rank])
                    v=f'{value:.3f}'; shown.append(f'<strong>{v}</strong>' if rank==1 else f'<u>{v}</u>' if rank==2 else v)
                rankings[ds][group][name]={'first':places.count(1),'second':places.count(2)}
                md.append('| '+label+' | '+' | '.join(shown)+' |')
                body.append('<tr><th>'+html.escape(label)+'</th>'+''.join('<td>'+x+'</td>' for x in shown)+'</tr>')
            sheet.freeze_panes='B2'; sheet.column_dimensions['A'].width=32
            sections.append('<h2>'+title+'</h2><div class="scroll"><table><tr>'+''.join('<th>'+html.escape(h)+'</th>' for h in headers)+'</tr>'+''.join(body)+'</table></div>')
            md.append('')
    path=out/f'{split}_seed32_tables.xlsx'; wb.save(path); reread=load_workbook(path)
    for title,ri,ci,value,rank in checks:
        cell=reread[title].cell(ri,ci)
        if rank is None: assert cell.value==value
        else:
            assert abs(cell.value-value)<1e-12
            assert cell.font.bold==(rank==1) and cell.font.underline==('single' if rank==2 else None)
    assert len(cells)==660
    (out/f'{split}_tables.md').write_text('\n'.join(md)+'\n')
    (out/f'{split}_tables.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><style>body{font-family:system-ui;margin:24px}.scroll{overflow:auto}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:7px;white-space:nowrap;text-align:right}</style>'+''.join('<p>'+html.escape(n)+'</p>' for n in notes)+''.join(sections)+'</html>')
    with (out/f'{split}_cells.csv').open('w',newline='') as f:
        writer=csv.writer(f); writer.writerow(['dataset','model','table','condition','metric','value','rank']); writer.writerows(cells)
    write(out/'results_and_provenance.json',dict(plan_sha256=plan['plan_sha256'],split=split,records=records,rankings=rankings,
          collector_sha256=sha(Path(__file__)),candidate='patch16',original_tables_preserved=True))
    write(out/'verification.json',dict(cells=660,model_labels=40,formats_verified=True,counts_verified=True,
          candidate='patch16',split=split,plan_sha256=plan['plan_sha256'],unchanged_records=19))
    return out, rankings


def main():
    p=check(); root=output(p)
    assert read(root/'TEST_COMPLETE.json')['plan_sha256']==p['plan_sha256']
    assert read(root/'TEST_COMPLETE.json')['count']==1
    selected={s:read(root/(s+'.json')) for s in ('validation','test')}
    for s,v in selected.items():
        assert v['passed'] and v['split']==s and v['plan_sha256']==p['plan_sha256']
        assert v['checkpoint_sha256']==p['checkpoint_sha256'] and v['best_epoch']==p['selected_epoch']
    assert selected['test']['teacher_forcing_positions']==0
    before={s:read(ROOT/f'paper/{s}_seed32/results_and_provenance.json') for s in selected}
    directories={}; all_ranks={}
    for split, source in before.items():
        records={ds:{('ALFNet' if name=='Focal' else name):copy.deepcopy(v) for name,v in items.items()} for ds,items in source['records'].items()}
        old=copy.deepcopy(records)
        # Replace the whole record: never retain the old checkpoint's identity.
        record=copy.deepcopy(selected[split])
        record.update(dataset='PEMSBAY',model='ALFNet',seed=32,selected_variant='patch16',
                      checkpoint=p['checkpoint'],config=p['config'],config_sha256=p['config_sha256'],
                      training_plan_sha256=p['training_plan_sha256'])
        records['PEMSBAY']['ALFNet']=record
        for ds,items in records.items():
            for name,v in items.items():
                if (ds,name)!=('PEMSBAY','ALFNet'): assert v==old[ds][name]
                for group,conditions in GROUPS.items():
                    for c in conditions:
                        for field,val in old[ds][name]['table'][group][c].items():
                            if field.endswith('_count'): assert v['table'][group][c][field]==val
        if split=='test':
            assert record['scaler']==old['PEMSBAY']['ALFNet']['scaler']
            assert record['test_exposure_hash']==old['PEMSBAY']['ALFNet']['test_exposure_hash']
        directory,ranks=publish(records,split,p); directories[split]=str(directory.relative_to(ROOT)); all_ranks[split]=ranks
    check()  # frozen original source records and checkpoint still unchanged
    write(root/'TABLES_COMPLETE.json',dict(plan_sha256=p['plan_sha256'],directories=directories,
          collector='scripts/collect_patch16_update.py',collector_sha256=sha(Path(__file__)),no_new_inference=True))
    status=read(root/'queue_status.json'); write(root/'queue_status_before_table_repair.json',status)
    status.update(test='complete',tables='complete',table_repair='offline collector; original queue log preserved')
    write(root/'queue_status.json',status)
    for split,ranks in all_ranks.items(): print(split,{ds:{g:r['ALFNet'] for g,r in groups.items()} for ds,groups in ranks.items()})


if __name__=='__main__': main()
