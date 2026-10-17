"""Authorized seed32 completion: 29 new runs, 26 reused rows, validation only."""
import argparse
import copy
import csv
import fcntl
import math
import os
from pathlib import Path
import subprocess
import sys
import time
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.run_metrla_hparams import read, write, sha, digest, identity

NAME = 'baseline_extension_20260927'
CONFIGS = ROOT / 'configs/paper' / NAME
PLAN = CONFIGS / 'plan.json'
REUSE = ROOT / 'paper/validation_seed32_flow_selected_20260927/results_and_provenance.json'
MODELS = ['UIGNN', 'GinAR', 'STFiT', 'SLPF', 'STGCN', 'STID', 'PatchSTG', 'STBP']
DATASETS = ['METRLA', 'PEMSBAY', 'PEMS03', 'PEMS04', 'PEMS08']
CAPACITY = dict(PEMS03=12, PEMS04=10, PEMS08=6)


def output(p):
    return ROOT / 'artifacts/experiments' / NAME / p['plan_sha256']


def freeze():
    assert not PLAN.exists(), 'frozen identity exists'
    CONFIGS.mkdir(parents=True, exist_ok=True)
    configs, audit, sources = {}, [], [REUSE, ROOT/'paper'/NAME/'PLAN.md']
    for ds in DATASETS:
        models = ['Merlin'] + (MODELS if ds in CAPACITY else [])
        for m in models:
            source = ROOT/f'artifacts/models/main/METRLA/{"STID" if m == "Merlin" else m}/seed32/config.yaml'
            base = yaml.safe_load(source.read_text())
            data_source = ROOT/f'configs/ALFNet/{ds}.yaml'
            base['dataset'] = copy.deepcopy(yaml.safe_load(data_source.read_text())['dataset'])
            base['runtime'].update(model=m, model_variant=f'{NAME}_{ds}_{m}')
            base['training']['num_workers'] = 0
            if m == 'PatchSTG':
                base['family']['model']['spa_patchsize'] = CAPACITY[ds]
            if m == 'Merlin':
                base['family'] = {'model': {}}
                base['training'].update(lr=.0002, weight_decay=.0001, batch_size=16,
                    max_epoch=100, patience=10, scheduler='None', loss='mae',
                    early_stopping_start_epoch=0)
            key = f'{ds}_{m}'
            path = CONFIGS/(key+'.yaml')
            path.write_text(yaml.safe_dump(base, sort_keys=False))
            configs[key] = dict(path=str(path.relative_to(ROOT)), sha256=sha(path), dataset=ds, model=m, seed=32)
            sources.extend([source, data_source])
            historical = []
            for r in sorted((ROOT/f'artifacts/models/history/{ds}/{m}/seed32').glob('*/run_report.json')):
                report = read(r)
                assert report['early_stopping_metric'] != 'validation 33-cell main/LifeAvg/MAE', ('audit needs review', r)
                historical.append(dict(report=str(r.relative_to(ROOT)), sha256=sha(r),
                    selection=report['early_stopping_metric'], old_score=report['best_validation_life_avg_mae'],
                    canonical_score=read(r.with_name('development_validation_table.json'))['main']['LifeAvg']['MAE']))
                sources.append(r)
            audit.append(dict(key=key, action='train', reason='absent' if not historical else 'old noncanonical checkpoint selection', historical=historical))
        manifest = ROOT/f'configs/paper/v4_data_manifests/{ds}_run_manifest.json'
        sources.append(manifest)
        d = yaml.safe_load((ROOT/f'configs/ALFNet/{ds}.yaml').read_text())['dataset']
        sources.extend(ROOT/d[k] for k in ('data_path', 'unknown_nodes_path'))
    sources.append(ROOT/'args.py')
    for folder in ('model', 'layers', 'lib', 'experiments', 'utils', 'data_provider', 'scripts'):
        sources.extend((ROOT/folder).rglob('*.py'))
    sources.extend(CONFIGS.glob('*.yaml'))
    p = dict(kind=NAME, status='authorized_validation_only', configs=configs, seed=32,
        max_concurrent=5, test_authorized=False, reuse=str(REUSE.relative_to(ROOT)), audit=audit,
        selection='full validation 33-cell LifeAvg MAE; earlier epoch on exact tie',
        merlin='teacher/student each max100 patience10; canonical validation selector; train-K privileged teacher only',
        files={str(f.relative_to(ROOT)): sha(f) for f in sorted(set(sources))},
        git_head=subprocess.check_output(['git','rev-parse','HEAD'], cwd=ROOT,text=True).strip())
    p['plan_sha256'] = digest(p)
    write(PLAN, p)
    print(p['plan_sha256'])


def check():
    p = read(PLAN)
    assert p['plan_sha256'] == digest(p) and len(p['configs']) == 29
    assert not p['test_authorized'] and p['max_concurrent'] == 5
    for f, expected in p['files'].items():
        assert sha(ROOT/f) == expected, f
    return p


def preflight():
    import torch
    from experiments.v4_runner import build_v4_session
    from scripts.train_v4 import train
    from scripts.train_merlin_v4 import train as merlin
    torch.set_num_threads(4)
    p = check()
    root = output(p)
    records = read(ROOT/p['reuse'])['records']
    identities = {}
    for ds in DATASETS:
        if ds in CAPACITY:
            evidence = ROOT/records[ds]['MoGERNN']['reuse_verification']
            identities[ds] = read(evidence)['identity']
        else:
            cp = torch.load(ROOT/f'artifacts/models/main/{ds}/MoGERNN/seed32/best.pt',map_location='cpu',weights_only=False)
            identities[ds] = {k: cp[k] for k in ('scaler','seed_contract','train_exposure_hash','validation_exposure_hash','static_asset_hash')}
            del cp
    results = []
    for key, row in p['configs'].items():
        saved = root/'preflight'/(key+'.json')
        if saved.exists():
            r = read(saved)
            assert r['passed'] and r['plan_sha256'] == p['plan_sha256']
            results.append(r)
            continue
        print('PREFLIGHT_START', key, flush=True)
        cfg = ROOT/row['path']
        s = build_v4_session(cfg,device='cpu',run_seed=32,num_workers=0)
        assert identity(s) == identities[row['dataset']], (key, 'protocol mismatch')
        assert not s.args.lifecycle_role_features and not s.args.lifecycle_history_age_feature
        if row['model'] == 'PatchSTG':
            capacity = CAPACITY[row['dataset']]
            assert len(s.train_context.layout_parts) == 32
            assert max(map(len, s.train_context.layout_parts)) <= capacity
        del s
        torch.cuda.reset_peak_memory_stats()
        started = time.time()
        if row['model'] == 'Merlin':
            report = merlin(cfg, root/'smoke'/key, p, smoke=True)
            teacher_report = read(report['teacher_report'])
            assert all(e[x]['skipped'] == 0 for e in teacher_report['history'] for x in ('train','validation'))
        else:
            report = train(cfg,None,device='cuda:0',seed=32,num_workers=0,plan=p,save=False,
                           max_epochs=1,max_train_batches=2,max_validation_batches=2)
        assert all(e[x]['skipped'] == 0 for e in report['history'] for x in ('train','validation'))
        r = dict(key=key,passed=True,plan_sha256=p['plan_sha256'],seconds=time.time()-started,
                 peak_reserved_gib=torch.cuda.max_memory_reserved()/1024**3)
        write(saved,r)
        results.append(r)
        torch.cuda.empty_cache()
        print('PREFLIGHT_OK',key,r,flush=True)
    write(root/'PREFLIGHT_COMPLETE.json',dict(passed=True,plan_sha256=p['plan_sha256'],reports=results))


def worker(key):
    import torch
    from scripts.train_v4 import train
    from scripts.train_merlin_v4 import train as merlin
    torch.set_num_threads(4)
    p = check()
    root = output(p)
    assert read(root/'PREFLIGHT_COMPLETE.json')['plan_sha256'] == p['plan_sha256']
    row = p['configs'][key]
    if row['model'] == 'Merlin':
        report = merlin(ROOT/row['path'],root/'runs'/key,p)
    else:
        report = train(ROOT/row['path'],root/'runs'/key,device='cuda:0',seed=32,num_workers=0,
                       plan=p,collect_development_table=True)
    assert all(e[x]['skipped'] == 0 for e in report['history'] for x in ('train','validation'))
    write(root/(key+'_complete.json'),dict(passed=True,plan_sha256=p['plan_sha256'],
        report_path=str(Path(report['run_directory'])/'run_report.json')))


def collect():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    p = check()
    root = output(p)
    records = copy.deepcopy(read(ROOT/p['reuse'])['records'])
    for key,row in p['configs'].items():
        done = read(root/(key+'_complete.json'))
        assert done['passed'] and done['plan_sha256'] == p['plan_sha256']
        rp = Path(done['report_path'])
        report = read(rp)
        assert report['plan_sha256'] == p['plan_sha256'] and report['config_sha256'] == row['sha256']
        assert min((e['validation']['checkpoint_life_avg_mae'],e['epoch']) for e in report['history'])[1] == report['best_epoch']
        assert all(e[x]['skipped'] == 0 for e in report['history'] for x in ('train','validation'))
        assert rp.with_name('COMPLETE').exists()
        records[row['dataset']][row['model']] = dict(table=report['development_validation_table'],
            report=str(rp.relative_to(ROOT)),report_sha256=sha(rp),checkpoint=str(rp.with_name('best.pt').relative_to(ROOT)),
            checkpoint_sha256=sha(rp.with_name('best.pt')),config=row['path'],config_sha256=row['sha256'],best_epoch=report['best_epoch'])
    out = ROOT/'paper'/('validation_seed32_'+NAME)
    out.mkdir(parents=True,exist_ok=True)
    groups = dict(main=['Unseen','Steady','Join','Fault','LifeAvg'],phase=['Cold','Warm','Ready','Onset','Decay','Outage'])
    names = ['ALFNet','UIGNN','GinAR','Merlin','MoGERNN','STFiT','SLPF','STGCN','STID','PatchSTG','STBP']
    md = ['# 五数据集 seed32 validation 主表与阶段表','','55条模型记录；第一名加粗，第二个不同值下划线；按未舍入值排序。',
          '本轮未打开 test；MAE/RMSE 不跨 flow/speed 汇总。SLPF 为单阶段预测适配，STBP 为离线结构编码适配。','']
    wb = Workbook()
    wb.remove(wb.active)
    cells,checks = [],[]
    for ds in DATASETS:
        assert set(records[ds]) == set(names)
        for group,conditions in groups.items():
            keys = [(c,m) for c in conditions for m in ('MAE','RMSE','MAPE')]
            title = ds+'_'+group
            sheet = wb.create_sheet(title)
            header = ['Model']+[c+'/'+m for c,m in keys]
            sheet.append(header)
            md += ['## '+title,'','| '+' | '.join(header)+' |','|'+'|'.join(['---']+['---:']*len(keys))+'|']
            for ri,name in enumerate(names,2):
                shown = []
                sheet.cell(ri,1,name)
                for ci,(c,m) in enumerate(keys,2):
                    val = records[ds][name]['table'][group][c][m]
                    assert math.isfinite(val)
                    reference = records[ds]['ALFNet']['table'][group][c]
                    for field,count in reference.items():
                        if field.endswith('_count'):
                            assert records[ds][name]['table'][group][c][field] == count,(ds,name,c,field)
                    rank = sorted(set(records[ds][n]['table'][group][c][m] for n in names)).index(val)+1
                    cell = sheet.cell(ri,ci,val)
                    cell.number_format = '0.000'
                    cell.font = Font(bold=rank==1,underline='single' if rank==2 else None)
                    text = f'{val:.3f}'
                    shown.append('**'+text+'**' if rank==1 else '<u>'+text+'</u>' if rank==2 else text)
                    cells.append([ds,name,group,c,m,val,rank])
                    checks.append((title,ri,ci,val,rank))
                md.append('| '+name+' | '+' | '.join(shown)+' |')
            md.append('')
    assert len(cells) == 1815
    wb.save(out/'final_tables.xlsx')
    reread = load_workbook(out/'final_tables.xlsx')
    for title,r,c,v,rank in checks:
        cell = reread[title].cell(r,c)
        assert abs(cell.value-v) < 1e-10 and cell.font.bold == (rank==1)
        assert cell.font.underline == ('single' if rank==2 else None)
    (out/'final_tables.md').write_text('\n'.join(md)+'\n')
    with (out/'cells.csv').open('w',newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['dataset','model','section','condition','metric','value','rank'])
        writer.writerows(cells)
    write(out/'results_and_provenance.json',dict(split='validation',seed=32,plan_sha256=p['plan_sha256'],records=records,test_accessed=False))
    write(out/'verification.json',dict(passed=True,cells=1815,models_per_dataset=11,excel_readback=True,test_accessed=False))
    write(root/'COMPLETE.json',dict(passed=True,table_directory=str(out.relative_to(ROOT))))


def queue():
    p = check()
    root = output(p)
    root.mkdir(parents=True,exist_ok=True)
    lock = (root/'queue.lock').open('a')
    fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
    assert not (root/'queue_status.json').exists(), 'inspect existing queue before restart'
    if not (root/'PREFLIGHT_COMPLETE.json').exists():
        preflight()
    gate = read(root/'PREFLIGHT_COMPLETE.json')
    peaks = {r['key']:r['peak_reserved_gib'] for r in gate['reports']}
    # Start Merlin across datasets first; each teacher/student pair uses one slot.
    order = sorted(p['configs'],key=lambda k:(p['configs'][k]['model']!='Merlin',k))
    jobs,active,last = [dict(key=k,status='queued') for k in order],{},0.
    env = dict(os.environ,OMP_NUM_THREADS='4',MKL_NUM_THREADS='4',PYTHONUNBUFFERED='1',MPLCONFIGDIR='/tmp/mogernn-mpl')
    while any(j['status'] in ('queued','running') for j in jobs):
        for key,(proc,log) in list(active.items()):
            code = proc.poll()
            if code is not None:
                job = next(j for j in jobs if j['key']==key)
                job.update(status='complete' if code==0 and (root/(key+'_complete.json')).exists() else 'failed',exit_code=code,finished=time.time())
                log.close()
                del active[key]
        if any(j['status']=='failed' for j in jobs):
            for j in jobs:
                if j['status']=='queued':
                    j['status']='blocked'
        pending = next((j for j in jobs if j['status']=='queued'),None)
        if pending and len(active)<5 and time.monotonic()-last>30:
            free = float(subprocess.check_output(['nvidia-smi','--query-gpu=memory.free','--format=csv,noheader,nounits','-i','0'],text=True).strip())/1024
            if free > max(6.,peaks[pending['key']]*1.5+3.):
                key = pending['key']
                log = (root/(key+'.log')).open('x')
                proc = subprocess.Popen([sys.executable,'-u',str(Path(__file__).resolve()),'worker','--key',key],cwd=ROOT,env=env,stdout=log,stderr=subprocess.STDOUT)
                pending.update(status='running',pid=proc.pid,started=time.time())
                active[key]=proc,log
                last=time.monotonic()
        write(root/'queue_status.json',dict(plan_sha256=p['plan_sha256'],jobs=jobs,updated=time.time()))
        time.sleep(15)
    assert all(j['status']=='complete' for j in jobs),'worker failure; no automatic retry'
    collect()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode',choices=['freeze','check','preflight','worker','queue','collect'])
    parser.add_argument('--key')
    args = parser.parse_args()
    if args.mode == 'worker':
        worker(args.key)
    elif args.mode == 'check':
        print(check()['plan_sha256'])
    else:
        globals()[args.mode]()
