"""Build paper-ready, fixed-anchor OFAT tables from saved validation results."""
import copy
import csv
import json
from pathlib import Path
import shutil
import sys

import yaml
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from scripts import run_hparam_completion as study
from scripts.collect_hparams_final import params, GROUPS, METRICS, OUT, SELECTED

ANCHORS=dict(METRLA='METRLA_baseline_reused',PEMSBAY='PEMSBAY_patch16',
             PEMS03='PEMS03_capacity_high',PEMS04='PEMS04_lr_half',PEMS08='PEMS08_lr_half')
INITIAL=dict(METRLA='METRLA_baseline_reused',PEMSBAY='PEMSBAY_baseline_reused',
             PEMS03='PEMS03_default',PEMS04='PEMS04_default',PEMS08='PEMS08_default')
AXES=dict(capacity='Patch容量',lr_multiplier='学习率倍率',dropout='Dropout',weight_decay='Weight decay')
ALLOWED={
    '/runtime/model_variant', '/training/lr', '/training/weight_decay',
    '/family/recipe/graph_expert_lr','/family/recipe/patch_expert_lr','/family/recipe/gate_lr',
    '/family/experts/patch/resolved/family/model/spa_patchsize',
    '/family/experts/patch/resolved/family/model/crossformer_dropout',
}


def flatten(x,p=''):
    if not isinstance(x,dict):return {p:x}
    out={}
    for k,v in x.items():out.update(flatten(v,p+'/'+str(k)))
    return out


def expanded(row):
    cfg=yaml.safe_load((ROOT/row['config']).read_text())
    for e in ('graph','patch'):
        x=cfg['family']['experts'][e]
        x['resolved']=yaml.safe_load((ROOT/x.pop('config_path')).read_text())
    return flatten(cfg)


def signature(ws):
    return [(c.coordinate,c.value,c.number_format,c.font.bold,c.font.underline)
            for row in ws for c in row if c.value is not None]


def main():
    study.check()
    rows=study.read(ROOT/'paper/hparams_completion_20260928/results.json')['rows']
    source=OUT/'final_tables.xlsx';wb=load_workbook(source)
    protected={n:signature(wb[n]) for n in wb.sheetnames if n.endswith(('_main','_phase')) or n=='ALFNet超参数'}
    backup=OUT/'final_tables_before_ofat_20260930.xlsx'
    if not backup.exists():shutil.copy2(source,backup)
    if '全部超参数实验' in wb.sheetnames:wb['全部超参数实验'].title='55项配置索引'
    for name in list(wb.sheetnames):
        if name.startswith('HP_') or name=='超参数实验说明':del wb[name]
    intro=wb.create_sheet('超参数实验说明',wb.sheetnames.index('55项配置索引'))
    instructions=[
        '论文数据：每个数据集固定同一参考点，四组单变量实验分别给出完整主表/阶段表；后半部分列出相对该参考点的已跑组合。',
        '参考点固定为2026-09-28补齐实验时的锚点，METRLA沿用原Focal；不是2026-09-30最终选定配置，不能标为围绕最终配置的完整四轴实验。',
        '同seed32、train/validation数据身份、Focal模型、验证LifeAvg MAE选checkpoint；max100/patience10/start40/batch64。',
        '每条曲线包含同一个参考点，重复展示只是复用结果而非重复训练。不同参数组的最小单元格不能拼接成一个模型。',
        '组合定义依赖参考点：表末额外标明相对最早默认配置改变的参数。逐次补点形成的单变量实验也可能是相对最早默认的多参数组合。',
        '超参数表不沿用对比算法排名样式：完整显示原始数值供论文作图，不根据本表再次宣称显著性。旧55项紧凑表保留为配置索引。',
        '无新训练或test；容量是token容量，分区固定32。LR倍率同时作用于三个参数组，Dropout为crossformer_dropout，Weight decay为AdamW参数。',
    ]
    for text in instructions:intro.append([text])
    intro.append(['数据集','固定参考run','容量','LR倍率','Dropout','Weight decay','单变量/参考配置数','额外组合数'])
    long=[];index=[];audit={};md=['# 论文用超参数实验：固定参考点的单变量结果','',*instructions,'']
    covered=set();single_total=combo_total=0
    for ds,anchor in ANCHORS.items():
        keys=[k for k,r in rows.items() if r['dataset']==ds]
        pp={k:params(rows[k]) for k in keys};anchor_cfg=expanded(rows[anchor]);ap=pp[anchor]
        changes={k:[a for a in AXES if pp[k][a]!=ap[a]] for k in keys}
        initial=pp[INITIAL[ds]]
        original_changes={k:[a for a in AXES if pp[k][a]!=initial[a]] for k in keys}
        for k in keys:
            cfg=expanded(rows[k]);diff={f for f in cfg.keys()|anchor_cfg.keys() if cfg.get(f)!=anchor_cfg.get(f)}
            assert diff<=ALLOWED,(k,diff-ALLOWED)
            assert cfg['/training/seed']==32
            rp=study.read(ROOT/rows[k]['report'])
            assert rp['early_stopping_metric']=='validation 33-cell main/LifeAvg/MAE'
            # Compare saved train/validation identity without loading any dataset.
            import torch
            state=torch.load(ROOT/rows[k]['checkpoint'],map_location='cpu',weights_only=False)
            ident={f:state[f] for f in ('payload_sha256','unseen_sha256','static_asset_hash','train_exposure_hash','validation_exposure_hash','scaler','seed_contract')}
            if k==keys[0]:reference_identity=ident
            assert ident==reference_identity,(ds,k)
            del state
            index.append(dict(dataset=ds,run=k,anchor=anchor,changed_axes=changes[k],
                changed_from_initial=original_changes[k],config=rows[k]['config'],checkpoint=rows[k]['checkpoint'],
                checkpoint_sha256=rows[k]['checkpoint_sha256'],report=rows[k]['report'],parameters=pp[k],
                selected_final=SELECTED[ds]==k))
        single=[k for k in keys if len(changes[k])<=1];combos=[k for k in keys if len(changes[k])>1]
        assert len(single)+len(combos)==len(keys)
        single_total+=len(single);combo_total+=len(combos)
        intro.append([ds,anchor,*[ap[a] for a in AXES],len(single),len(combos)])
        audit[ds]=dict(anchor=anchor,anchor_parameters=ap,single=single,combinations=combos,identity_equal=True)
        for g,conditions in GROUPS.items():
            name=f'HP_{ds}_'+('主表' if g=='main' else '阶段表')
            ws=wb.create_sheet(name,wb.sheetnames.index('55项配置索引'))
            ws.append([name+'：四组单变量实验 + 已跑组合'])
            ws.append([f'固定参考点 {anchor}; C={ap["capacity"]}, LR×{ap["lr_multiplier"]}, dropout={ap["dropout"]}, WD={ap["weight_decay"]}; seed32/validation'])
            metric_keys=[(c,m) for c in conditions for m in METRICS]
            header=['实验组','run','本轴取值','Patch容量','LR倍率','Dropout','Weight decay','参考点','最终选型']+[c+'/'+m for c,m in metric_keys]
            md += [f'## {name}','',f'参考点：`{anchor}`。','']
            blocks=[(AXES[a],a,sorted([k for k in keys if all(pp[k][b]==ap[b] for b in AXES if b!=a)],key=lambda k:pp[k][a])) for a in AXES]
            blocks.append(('已有组合（相对此参考点改变至少两轴）',None,combos))
            for label,axis,block in blocks:
                ws.append([label]);title_row=ws.max_row
                ws.cell(title_row,1).font=Font(bold=True,color='FFFFFF');ws.cell(title_row,1).fill=PatternFill('solid',fgColor='305496')
                ws.append(header)
                for c in ws[ws.max_row]:c.font=Font(bold=True)
                md += [f'### {label}','', '| '+' | '.join(header)+' |','|'+'|'.join(['---']*len(header))+'|']
                if not block:
                    ws.append(['无已完成组合']);md+=['无已完成组合。',''];continue
                for k in block:
                    covered.add(k);t=rows[k]['table'];v=pp[k]
                    values=[label,k,v[axis] if axis else '+'.join(AXES[a] for a in changes[k]),*[v[a] for a in AXES],k==anchor,k==SELECTED[ds]]+[t[g][c][m] for c,m in metric_keys]
                    ws.append(values)
                    for cell in list(ws[ws.max_row])[9:]:cell.number_format='0.000'
                    md.append('| '+' | '.join(f'{x:.4f}' if isinstance(x,float) else str(x) for x in values)+' |')
                    for c,m in metric_keys:long.append([ds,label,axis or 'combination',anchor,k,v[axis] if axis else '',*[v[a] for a in AXES],g,c,m,t[g][c][m]])
                ws.append([]);md.append('')
            ws.append(['组合追溯：以下改变轴相对最早默认配置计算，避免将不同参考点混淆。'])
            ws.append(['run','相对最早默认改变的参数','相对本表参考点改变的参数'])
            for k in keys:
                if len(original_changes[k])>1:
                    ws.append([k,'+'.join(AXES[a] for a in original_changes[k]),'+'.join(AXES[a] for a in changes[k]) or '本表参考点'])
            ws.freeze_panes='J5';ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=35
            for ci in range(3,len(header)+1):ws.column_dimensions[get_column_letter(ci)].width=17
    assert len(covered)==55 and single_total==47 and combo_total==8
    wb.save(source)
    loaded=load_workbook(source)
    for n,s in protected.items():assert signature(loaded[n])==s,n
    verified=0
    for ds in ANCHORS:
        for g,conditions in GROUPS.items():
            ws=loaded[f'HP_{ds}_'+('主表' if g=='main' else '阶段表')]
            for row in ws.iter_rows():
                key=row[1].value if len(row)>1 else None
                if key not in rows or row[0].value not in list(AXES.values())+['已有组合（相对此参考点改变至少两轴）']:continue
                vals=[rows[key]['table'][g][c][m] for c in conditions for m in METRICS]
                for cell,value in zip(row[9:],vals):assert abs(cell.value-value)<1e-10;verified+=1
    assert verified==len(long)==2310
    with (OUT/'paper_hparams_full_cells.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(['dataset','group','axis','anchor','run','axis_value',*AXES,'section','condition','metric','value']);w.writerows(long)
    study.write(OUT/'paper_hparams_registry.json',dict(anchors=ANCHORS,initial=INITIAL,runs=index,audit=audit,split='validation',seed=32))
    study.write(OUT/'paper_hparams_verification.json',dict(passed=True,unique_runs=55,single_or_anchor=47,additional_combinations=8,verified_cells=verified,previous_main_phase_hparams_sheets_unchanged=True,full_config_diff_verified=True,saved_train_validation_identity_verified=True,test_accessed=False))
    (OUT/'PAPER_HYPARAMETER_EXPERIMENTS.md').write_text('\n'.join(md)+'\n')
    print(json.dumps(dict(single_or_anchor=single_total,combinations=combo_total,verified_cells=verified,anchors=ANCHORS),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
