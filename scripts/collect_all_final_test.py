"""Publish separate validation-selected and exploratory-test-selected tables."""
import csv,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts import test_all_final_20261001 as run
from scripts.run_metrla_hparams import read,write
from scripts.organize_hparam_ofat import ANCHORS,AXES
from scripts.collect_hparams_final import params,GROUPS,METRICS

def choose(rows,keys,baseline):
    def vec(k):return [rows[k]['table']['main']['LifeAvg'][m] for m in METRICS]
    ref=[min(r['table']['main']['LifeAvg'][m] for r in baseline.values()) for m in METRICS]
    wins={k:sum(rows[k]['table'][g][c][m]<min(r['table'][g][c][m] for r in baseline.values()) for g,cs in GROUPS.items() for c in cs for m in METRICS) for k in keys}
    eligible=[k for k in keys if all(a<b for a,b in zip(vec(k),ref))]
    if not eligible:eligible=[k for k in keys if not any(all(a<=b for a,b in zip(vec(j),vec(k))) and any(a<b for a,b in zip(vec(j),vec(k))) for j in keys if j!=k)]
    return min(eligible,key=lambda k:(-wins[k],sum(a/b for a,b in zip(vec(k),ref))/3,vec(k)[0],k)),wins

def publish_tables(out,tag,records,entries,selection):
    wb=Workbook();wb.active.title='说明';wb.active.append([tag]);wb.active.append(['按未舍入数值排名，第一加粗、第二下划线。单seed；test选型版为探索性结果，不是独立确认。'])
    md=['# '+tag,'','Validation选型和test探索选型分别保存；不能混称未参与选择的独立test。',''];cells=[];checks=[]
    for ds,models in records.items():
        for g,cs in GROUPS.items():
            ws=wb.create_sheet(ds+'_'+g);keys=[(c,m) for c in cs for m in METRICS];h=['模型']+[c+'/'+m for c,m in keys];ws.append(h)
            md += ['## '+ds+' '+g,'','| '+' | '.join(h)+' |','|'+'|'.join(['---']+['---:']*len(keys))+'|']
            for ri,(n,r) in enumerate(models.items(),2):
                ws.cell(ri,1,n);shown=[]
                for ci,(c,m) in enumerate(keys,2):
                    v=r['table'][g][c][m];rank=sorted({x['table'][g][c][m] for x in models.values()}).index(v)+1
                    cell=ws.cell(ri,ci,v);cell.number_format='0.000';cell.font=Font(bold=rank==1,underline='single' if rank==2 else None)
                    t=f'{v:.3f}';shown.append(f'**{t}**' if rank==1 else f'<u>{t}</u>' if rank==2 else t)
                    cells.append([ds,n,g,c,m,v,rank]);checks.append((ws.title,ri,ci,v,rank))
                md.append('| '+n+' | '+' | '.join(shown)+' |')
            ws.freeze_panes='B2';md.append('')
    hp=wb.create_sheet('ALFNet配置');hp.append(['dataset','run',*AXES,'config','checkpoint'])
    for ds,k in selection.items():
        q=params(entries[k]);hp.append([ds,k,*[q[a] for a in AXES],entries[k]['config'],entries[k]['checkpoint']])
    wb.save(out/(tag+'.xlsx'));(out/(tag+'.md')).write_text('\n'.join(md)+'\n')
    check=load_workbook(out/(tag+'.xlsx'))
    for s,i,j,v,r in checks:
        c=check[s].cell(i,j);assert abs(c.value-v)<1e-10 and c.font.bold==(r==1) and c.font.underline==('single' if r==2 else None)
    with (out/(tag+'_cells.csv')).open('w',newline='') as f:w=csv.writer(f);w.writerow(['dataset','model','section','condition','metric','value','rank']);w.writerows(cells)
    return len(checks)

def collect():
    p=run.check();root=run.output(p);out=run.OUT;out.mkdir(parents=True,exist_ok=True)
    rows={k:read(root/(k+'.json')) for k in p['entries']}
    assert len(rows)==117 and all(r['passed'] and r['split']=='test' and r['checkpoint_sha256']==p['entries'][k]['checkpoint_sha256'] for k,r in rows.items())
    selections={};validation={};exploratory={};wins={}
    for ds in run.VALID_SELECTED:
        dsrows={k:r for k,r in rows.items() if p['entries'][k]['dataset']==ds};first=next(iter(dsrows.values()))
        for r in dsrows.values():
            assert r['scaler']==first['scaler'] and r['test_exposure_hash']==first['test_exposure_hash']
            for g,cs in GROUPS.items():
                for c in cs:
                    for f,v in first['table'][g][c].items():
                        if f.endswith('_count'):assert r['table'][g][c][f]==v
        baseline={p['entries'][k]['model']:r for k,r in dsrows.items() if p['entries'][k]['model']!='ALFNet'}
        keys=[k for k in dsrows if p['entries'][k]['model']=='ALFNet'];selected,w=choose(rows,keys,baseline);wins.update(w);selections[ds]=selected
        validation[ds]=dict(ALFNet=rows[p['validation_selected'][ds]],**baseline)
        exploratory[ds]=dict(ALFNet=rows[selected],**baseline)
    count=publish_tables(out,'validation_selected_test',validation,p['entries'],p['validation_selected'])
    count+=publish_tables(out,'exploratory_test_selected',exploratory,p['entries'],selections)
    wb=Workbook();wb.active.title='说明';wb.active.append(['所有67项ALFNet test；按固定参考点四轴分组，其余组合列在下方。test曲线为探索分析，论文敏感性主分析仍用validation。'])
    allcells=[]
    for ds,anchor in ANCHORS.items():
        keys=[k for k,e in p['entries'].items() if e['dataset']==ds and e['model']=='ALFNet'];ps={k:params(p['entries'][k]) for k in keys};ap=ps[anchor]
        groups=[(a,sorted([k for k in keys if all(ps[k][b]==ap[b] for b in AXES if b!=a)],key=lambda k:ps[k][a])) for a in AXES]
        groups.append(('combinations',[k for k in keys if sum(ps[k][a]!=ap[a] for a in AXES)>1]))
        for g,cs in GROUPS.items():
            ws=wb.create_sheet(ds+'_'+g)
            for label,ks in groups:
                ws.append([label,'anchor='+anchor]);ws.append(['run',*AXES,'第一格/33']+[c+'/'+m for c in cs for m in METRICS])
                for k in ks:
                    ws.append([k,*[ps[k][a] for a in AXES],wins[k]]+[rows[k]['table'][g][c][m] for c in cs for m in METRICS])
                    for c in cs:
                        for m in METRICS:allcells.append([ds,label,k,g,c,m,rows[k]['table'][g][c][m]])
                ws.append([])
    wb.save(out/'all_hyperparameter_tests.xlsx')
    with (out/'all_hyperparameter_test_cells.csv').open('w',newline='') as f:w=csv.writer(f);w.writerow(['dataset','group','run','section','condition','metric','value']);w.writerows(allcells)
    write(out/'all_results_and_provenance.json',dict(plan_sha256=p['plan_sha256'],rows=rows,entries=p['entries'],validation_selected=p['validation_selected'],exploratory_test_selected=selections,first_cells=wins,test_based_selection=True))
    write(out/'verification.json',dict(passed=True,evaluations=117,comparison_table_cells=count,all_datasets_test_counts_match=True,test_based_selection_disclosed=True))
    write(root/'COMPLETE.json',dict(passed=True,evaluations=117))

if __name__=='__main__':collect()
