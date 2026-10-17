"""All five fixed WD=.001 ablations: full validation replay before test."""
import argparse,copy,csv,sys
from pathlib import Path
from contextlib import ExitStack
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.run_metrla_hparams import read,write,sha,digest
from scripts import run_pems08_wd_ablations_20261012 as train
from scripts import test_all_ablations_20261010 as base
from scripts import evaluate_speed_support_20261006 as runner
NAME='pems08_wd_ablation_test_20261012';OUT=ROOT/'paper'/NAME;PLAN=ROOT/'configs/paper'/NAME/'plan.json'
def output(p):return ROOT/'artifacts/evaluation'/NAME/p['plan_sha256']
def freeze():
    import torch,yaml
    assert not PLAN.exists();tp=train.check();root=train.output(tp)
    assert read(root/'COMPLETE.json')['passed']
    assert all(j['status']=='complete' and j['exit_code']==0 for j in read(root/'queue_status.json')['jobs'])
    source=train.OUT/'results.json';rows=read(source)['rows'];entries={};audits={};files=[source,train.PLAN,Path(__file__)]
    for k,e in tp['configs'].items():
        r=copy.deepcopy(rows[k]);cp=torch.load(ROOT/r['checkpoint'],map_location='cpu',weights_only=False);report=read(ROOT/r['report'])
        assert cp['seed']==32 and cp['plan_sha256']==tp['plan_sha256'] and cp['config_sha256']==sha(ROOT/r['config'])
        assert sha(ROOT/r['checkpoint'])==r['checkpoint_sha256'] and report['development_validation_table']==r['table']
        assert cp['epoch']==report['best_epoch']==min((h['validation']['checkpoint_life_avg_mae'],h['epoch']) for h in report['history'])[1]
        assert all(h[s]['skipped']==0 for h in report['history'] for s in ('train','validation'))
        assert all(torch.isfinite(v).all() for v in cp['model_state'].values() if v.is_floating_point())
        r.update(model='ALFNet',ablation=e['ablation'],epoch=cp['epoch'],training_config_sha256=cp['config_sha256'],training_plan_sha256=cp['plan_sha256']);entries[k]=r
        audits[k]=dict(best_epoch=cp['epoch'],epochs_run=report['epochs_run'],finite_weights=True,zero_skips=True)
        files.extend(ROOT/r[f] for f in ('config','checkpoint','report'))
        cfg=yaml.safe_load((ROOT/r['config']).read_text())
        files.extend(ROOT/cfg['family']['experts'][x]['config_path'] for x in ('graph','patch'))
        files.extend(ROOT/cfg['dataset'][f] for f in ('data_path','unknown_nodes_path'))
    fullplan=read(ROOT/'configs/paper/pems08_wd_test_20261012/plan.json')
    fullpath=ROOT/'artifacts/evaluation/pems08_wd_test_20261012'/fullplan['plan_sha256']/'13_PEMS08_weight_decay_0.001.json'
    full=read(fullpath);assert full['passed'] and full['checkpoint_sha256']==rows['PEMS08_Full']['checkpoint_sha256']
    files.extend([fullpath,ROOT/'paper/all_ablations_test_20261010/results_and_audit.json',ROOT/'paper/ablation_repeat_test_20261012/results.json'])
    for folder in ('model','layers','lib','experiments','data_provider','utils'):files.extend((ROOT/folder).rglob('*.py'))
    files.extend(ROOT/f for f in ('scripts/test_all_ablations_20261010.py','scripts/test_all_final_20261001.py','scripts/run_final_ablations_20261006.py','scripts/run_mechanism_ablations_20261008.py','scripts/run_reconstruction_gradient_20261009.py','scripts/train_v4.py','scripts/eval_release.py','scripts/evaluate_speed_support_20261006.py','args.py'))
    p=dict(kind=NAME,status='authorized_test',test_authorized=True,authorization='User explicitly asks to continue after training and determine whether test meets requirements.',entries=entries,reused_full=full,training_audit=audits,max_concurrent=5,reporting='All five validation-best checkpoints, no filtering or retraining; strict positive change in each LifeAvg metric relative to same-WD Full.',files={str(f.relative_to(ROOT)):sha(f) for f in set(files)})
    p['plan_sha256']=digest(p);write(PLAN,p);print(p['plan_sha256'])
def check():
    p=read(PLAN);assert digest(p)==p['plan_sha256'] and p['test_authorized'] and len(p['entries'])==5
    for f,h in p['files'].items():assert sha(ROOT/f)==h,f
    return p
def worker(k):
    with patch.object(base,'check',check),patch.object(base,'output',output):base.worker(k)
def collect():
    from openpyxl import Workbook
    p=check();root=output(p);rows={'PEMS08_Full':dict(p['reused_full'],label='Full')};replays={};deltas={};cells=[]
    for k,e in p['entries'].items():
        r=read(root/(k+'.json'));v=read(root/'validation'/(k+'.json'))
        assert r['passed'] and v['passed'] and r['checkpoint_sha256']==e['checkpoint_sha256']
        rows[k]=r;replays[k]=dict(passed=True,max_abs_difference=v['max_abs_difference'])
    old=read(ROOT/'paper/all_ablations_test_20261010/results_and_audit.json')['rows'];repeat=read(ROOT/'paper/ablation_repeat_test_20261012/results.json')['rows']['02_PEMS04_no_age']['repeat_ablation']
    metrics=('MAE','RMSE','MAPE');four={}
    for k,e in p['entries'].items():
        kind=e['ablation'];four[k]={}
        for ds in ('METRLA','PEMSBAY','PEMS04','PEMS08'):
            ref=rows['PEMS08_Full'] if ds=='PEMS08' else old[ds+'_Full']
            r=rows[k] if ds=='PEMS08' else next(v for n,v in old.items() if n.endswith('_'+ds+'_'+kind))
            if ds=='PEMS04' and kind=='no_age':r=repeat
            four[k][ds]={m:100*(r['table']['main']['LifeAvg'][m]/ref['table']['main']['LifeAvg'][m]-1) for m in metrics}
        deltas[k]=four[k]['PEMS08']
    for k,r in rows.items():
        for g in ('main','phase'):
            for c,v in r['table'][g].items():
                for f,n in v.items():
                    if f.endswith('_count'):assert n==rows['PEMS08_Full']['table'][g][c][f]
                for m in metrics:cells.append([k,r['label'],g,c,m,v[m]])
    q=[k for k,d in deltas.items() if all(x>0 for x in d.values())];q4=[k for k,d in four.items() if all(x>0 for a in d.values() for x in a.values())]
    write(OUT/'results.json',dict(rows=rows,validation_replays=replays,lifeavg_delta_percent=deltas,four_dataset_delta_percent=four,qualifying_pems08=q,qualifying_four_datasets=q4,plan_sha256=p['plan_sha256']))
    wb=Workbook();wb.remove(wb.active)
    for group in ('main','phase'):
        ws=wb.create_sheet(group);cols=[(c,m) for c in rows['PEMS08_Full']['table'][group] for m in metrics];ws.append(['Variant']+[c+'/'+m for c,m in cols])
        for r in rows.values():ws.append([r['label']]+[r['table'][group][c][m] for c,m in cols])
    ws=wb.create_sheet('四数据集LifeAvg变化');ws.append(['Variant']+[d+'/'+m+'%' for d in ('METRLA','PEMSBAY','PEMS04','PEMS08') for m in metrics]+['变差项/12'])
    for k,d in four.items():ws.append([p['entries'][k]['label']]+[d[ds][m] for ds in ('METRLA','PEMSBAY','PEMS04','PEMS08') for m in metrics]+[sum(x>0 for a in d.values() for x in a.values())])
    wb.save(OUT/'test_main_phase.xlsx')
    with (OUT/'test_cells.csv').open('w') as f:w=csv.writer(f);w.writerow(['run','label','section','condition','metric','value']);w.writerows(cells)
    lines=['# PEMS08 WD=.001 五项配对消融 test','','全部固定权重均已评估；Full复用相同checkpoint测试。PEMS04去Age使用用户保留的重复记录，其余三个数据集复用原有配对结果。',f'PEMS08全三项变差：{len(q)}/5；四数据集全部12项变差：{len(q4)}/5。','','| 变体 | MAE | RMSE | MAPE | 相对Full变化% | 四数据集变差项/12 |','|---|---:|---:|---:|---|---:|']
    for k,r in rows.items():
        a=r['table']['main']['LifeAvg'];ds=deltas.get(k);count=sum(x>0 for a in four[k].values() for x in a.values()) if k in four else 0
        lines.append('| '+r['label']+' | '+' | '.join(f'{a[m]:.6f}' for m in metrics)+' | '+('/'.join(f'{ds[m]:+.2f}' for m in metrics) if ds else '—')+f' | {count if ds else "—"} |')
    (OUT/'ANALYSIS.md').write_text('\n'.join(lines)+'\n');write(root/'COMPLETE.json',dict(passed=True,new_tests=5,reused_full=1,qualifying_pems08=q,qualifying_four_datasets=q4));print('\n'.join(lines))
def queue():
    p=check()
    with ExitStack() as s:
        for n,v in [('check',check),('output',output),('collect',collect),('KEYS',tuple(p['entries'])),('__file__',__file__)]:s.enter_context(patch.object(runner,n,v))
        runner.queue()
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['freeze','worker','queue','collect']);p.add_argument('--key');a=p.parse_args()
    if a.mode=='worker':worker(a.key)
    else:globals()[a.mode]()
