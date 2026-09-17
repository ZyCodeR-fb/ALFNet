"""Publish the locked release's complete test tables; no inference or selection."""
import csv,html,json,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.eval_release import check,output,read,write

MODELS=('ALFNet','UIGNN','GinAR','MoGERNN','STFiT','SLPF','STGCN','STID','PatchSTG','STBP')
GROUPS={'main':('Unseen','Steady','Join','Fault','LifeAvg'),'phase':('Cold','Warm','Ready','Onset','Decay','Outage')}
METRICS=('MAE','RMSE','MAPE')
LABELS={'ALFNet':'ALFNet（Focal）','STFiT':'ST-FiT','SLPF':'SLPF（单阶段）'}

def main():
    p=check();root=output(p);assert read(root/'TEST_COMPLETE.json')['count']==20
    records={ds:{} for ds in ('METRLA','PEMSBAY')}
    for row in p['runs']:
        r=read(root/'test'/(row['dataset']+'_'+row['model']+'.json'))
        assert r['passed'] and r['plan_sha256']==p['plan_sha256'] and r['seed']==32 and r['split']=='test'
        assert r['checkpoint_sha256']==row['checkpoint_sha256'] and r['teacher_forcing_positions']==0
        records[row['dataset']][row['model']]=r
    for ds,items in records.items():
        assert set(items)==set(MODELS);ref=items['ALFNet']
        for r in items.values():
            assert r['test_exposure_hash']==ref['test_exposure_hash'] and r['scaler']==ref['scaler']
            for group,rows in GROUPS.items():
                for label in rows:
                    for k,v in ref['table'][group][label].items():
                        if k.endswith('_count'):assert r['table'][group][label][k]==v
    out=ROOT/'paper/test_seed32';out.mkdir(exist_ok=True)
    notes=['范围：METR-LA、PEMS-BAY，seed32，冻结的Focal＋九基线，test结果。',
        '权重来自artifacts/models/main；18项通过历史33格validation一致性检查，2项ST-FiT通过固定随机种子的归档源码/当前源码配对全量validation，再打开test。',
        'ST-FiT保留源Gumbel随机评估；历史best未保存当时RNG，因此不声称逐值重播训练中的validation。',
        '使用checkpoint的训练scaler，test不拟合scaler，teacher forcing为0；不依据test重新选模型。',
        '每列按未舍入值排名，第一名加粗，第二个不同数值下划线；精确并列共享名次。',
        'SLPF是单阶段源预测主干适配；STBP为离线固定图主干；结构参数归纳适配继续披露。',
        'MAE/RMSE为原速度单位，MAPE为%；不是多种子均值，格数不代表独立统计检验。']
    wb=Workbook();ws=wb.active;ws.title='说明'
    for n in notes:ws.append([n])
    ws.column_dimensions['A'].width=130
    for row in ws:row[0].alignment=Alignment(wrap_text=True)
    md=['# 最终seed32 test结果','',*notes,''];sections=[];cells=[];verification=[];rankings={}
    for ds,items in records.items():
        rankings[ds]={}
        for group,rows in GROUPS.items():
            title=ds+('主表' if group=='main' else '阶段表');sheet=wb.create_sheet(title)
            keys=[(r,m) for r in rows for m in METRICS]
            sheet.append(['模型']+[r+'/'+m for r,m in keys]);body=[];rankings[ds][group]={}
            md+=['## '+title,'','| 模型 | '+' | '.join(r+'/'+m for r,m in keys)+' |','|'+'|'.join(['---']+['---:']*len(keys))+'|']
            for ri,name in enumerate(MODELS,2):
                sheet.cell(ri,1,LABELS.get(name,name));hs=[];ms=[];places=[]
                for ci,(r,m) in enumerate(keys,2):
                    value=items[name]['table'][group][r][m]
                    place=1+sorted(set(x['table'][group][r][m] for x in items.values())).index(value)
                    places.append(place);cell=sheet.cell(ri,ci,value);cell.number_format='0.000'
                    cell.font=Font(bold=place==1,underline='single' if place==2 else None)
                    verification.append((title,ri,ci,value,place));cells.append([ds,name,group,r,m,value,place])
                    v=f'{value:.3f}';hs.append('<td>'+(f'<strong>{v}</strong>' if place==1 else f'<u>{v}</u>' if place==2 else v)+'</td>')
                    ms.append(f'**{v}**' if place==1 else f'<u>{v}</u>' if place==2 else v)
                rankings[ds][group][name]={'first':places.count(1),'second':places.count(2)}
                md.append('| '+LABELS.get(name,name)+' | '+' | '.join(ms)+' |')
                body.append('<tr><th>'+html.escape(LABELS.get(name,name))+'</th>'+''.join(hs)+'</tr>')
            headers='<tr><th>模型</th>'+''.join(f'<th>{r}/{m}</th>' for r,m in keys)+'</tr>'
            sections.append('<h2>'+title+'</h2><div class="scroll"><table>'+headers+''.join(body)+'</table></div>');md.append('')
            sheet.freeze_panes='B2';sheet.column_dimensions['A'].width=24
            for ci in range(2,len(keys)+2):sheet.column_dimensions[get_column_letter(ci)].width=15
            for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183B56')
    path=out/'test_seed32_tables.xlsx';wb.save(path);reloaded=load_workbook(path)
    for name,ri,ci,v,k in verification:
        c=reloaded[name].cell(ri,ci);assert abs(c.value-v)<1e-12
        assert c.font.bold==(k==1) and c.font.underline==('single' if k==2 else None)
    assert len(verification)==660
    (out/'test_tables.md').write_text('\n'.join(md)+'\n')
    (out/'test_tables.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>最终test结果</title><style>body{font-family:system-ui;margin:28px}.scroll{overflow:auto}table{border-collapse:collapse}td,th{border:1px solid #ccd5dd;padding:8px;white-space:nowrap;text-align:right}u{text-underline-offset:3px}</style><h1>最终seed32 test结果</h1>'+''.join('<p>'+html.escape(n)+'</p>' for n in notes)+''.join(sections)+'</html>')
    with (out/'test_cells.csv').open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['dataset','model','table','condition','metric','value','rank']);writer.writerows(cells)
    write(out/'results_and_provenance.json',dict(plan_sha256=p['plan_sha256'],records=records,rankings=rankings))
    write(out/'verification.json',dict(runs=20,split='test',cells=660,formats_verified=True,shared_exposure_counts_scaler=True,teacher_forcing_positions=0))
    print(json.dumps({ds:{g:rankings[ds][g]['ALFNet'] for g in GROUPS} for ds in records},indent=2))

if __name__=='__main__':main()
