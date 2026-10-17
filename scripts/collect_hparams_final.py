"""Audit 55 ALFNet candidates and publish validation tables; no inference/test."""
import copy
import csv
import json
import math
from pathlib import Path
import sys

import torch
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import run_hparam_completion as study

OUT = ROOT / 'paper/validation_seed32_hparams_final_20260930'
METRICS = ('MAE', 'RMSE', 'MAPE')
GROUPS = dict(main=('Unseen','Steady','Join','Fault','LifeAvg'),
              phase=('Cold','Warm','Ready','Onset','Decay','Outage'))
SELECTED = dict(METRLA='METRLA_lr_double', PEMSBAY='23_PEMSBAY_lr_double',
                PEMS03='16_PEMS03_lr_half', PEMS04='PEMS04_lr_half', PEMS08='02_PEMS08_drop05')
REASONS = {
    'METRLA': '学习率×2同时具有最多第一格和最低LifeAvg三指标，明确保留。',
    'PEMSBAY': '容量16/学习率×2以LifeAvg小幅代价换取第一格11→18，三项LifeAvg仍领先全部基线；容量16/学习率×1保留为最低平均误差备选。',
    'PEMS03': '容量20/学习率×0.5与原配置同为6格第一；MAE小幅增加，RMSE与MAPE降低，选择更均衡的三指标结果。容量20/学习率×1保留为最低MAE备选。',
    'PEMS04': '容量10/学习率×0.5仍同时具有最多第一格和最低MAE/RMSE，维持原配置。',
    'PEMS08': '容量6/学习率×0.5/dropout0.05增加两格第一并改善MAE/RMSE，接受MAPE小幅增加。dropout0.2为最低MAE备选；旧学习率×2虽有11格第一，但三项LifeAvg均较所选差，不为额外两格接受该代价。',
}


def rank(models, name, g, c, m):
    return sorted({r['table'][g][c][m] for r in models.values()}).index(models[name]['table'][g][c][m])+1


def params(row):
    cfg = yaml.safe_load((ROOT / row['config']).read_text())
    patch = yaml.safe_load((ROOT / cfg['family']['experts']['patch']['config_path']).read_text())['family']['model']
    recipe = cfg['family']['recipe']
    lrs = [recipe[k] for k in ('graph_expert_lr','patch_expert_lr','gate_lr')]
    factors = [a/b for a,b in zip(lrs,(5e-4,1e-4,5e-5))]
    assert max(factors)-min(factors) < 1e-12
    return dict(capacity=patch['spa_patchsize'],lr_multiplier=factors[0],graph_lr=lrs[0],
                temporal_lr=lrs[1],gate_lr=lrs[2],dropout=patch['crossformer_dropout'],
                weight_decay=cfg['training']['weight_decay'],batch_size=cfg['training']['batch_size'],
                max_epoch=cfg['training']['max_epoch'],patience=cfg['training']['patience'],
                early_stopping_start_epoch=cfg['training']['early_stopping_start_epoch'],
                loss=cfg['training']['loss'])


def main():
    torch.set_num_threads(4)
    p = study.check(); root = study.output(p)
    jobs = study.read(root/'queue_status.json')['jobs']
    assert len(jobs)==24 and all(j['status']=='complete' and j['exit_code']==0 for j in jobs)
    assert study.read(root/'COMPLETE.json')['passed']
    source = ROOT/'paper/hparams_completion_20260928/results.json'
    data = study.read(source); rows = data['rows']
    assert len(rows)==55 and data['split']=='validation' and data['plan_sha256']==p['plan_sha256']
    base = study.read(study.BASE_TABLE)['records']; records = copy.deepcopy(base)
    audited = {}; candidates = []
    for key,row in rows.items():
        rp,cp = ROOT/row['report'],ROOT/row['checkpoint']
        report = study.read(rp)
        assert rp.with_name('COMPLETE').exists()
        assert report['development_validation_table']==row['table']
        assert report['seed']==32 and report['parameter_count']==1037529
        assert [e['epoch'] for e in report['history']]==list(range(report['epochs_run']))
        assert all(e[t]['skipped']==0 for e in report['history'] for t in ('train','validation'))
        assert min((e['validation']['checkpoint_life_avg_mae'],e['epoch']) for e in report['history'])[1]==report['best_epoch']
        config_sha = study.sha(ROOT/row['config'])
        if config_sha != report['config_sha256']:
            # Archived release configurations only relocate expert paths.
            provenance = study.read(rp.with_name('provenance.json'))
            assert config_sha==provenance['config_sha256']==row['config_sha256']
            assert study.sha(rp.with_name('config.original.yaml'))==report['config_sha256']==provenance['original_config_sha256']
            current=yaml.safe_load((ROOT/row['config']).read_text())
            original=yaml.safe_load(rp.with_name('config.original.yaml').read_text())
            for expert in ('graph','patch'):
                path=current['family']['experts'][expert]['config_path']
                assert yaml.safe_load((ROOT/path).read_text())==yaml.safe_load(rp.with_name(f'expert.{expert}.original.yaml').read_text())
                current['family']['experts'][expert]['config_path']=original['family']['experts'][expert]['config_path']
            assert current==original
        assert study.sha(cp)==row['checkpoint_sha256']
        state = torch.load(cp,map_location='cpu',weights_only=False)
        assert state['config_sha256']==report['config_sha256'] and state['plan_sha256']==report['plan_sha256']
        assert state['best_epoch']==report['best_epoch'] and state['seed']==32
        assert all(torch.isfinite(v).all() for v in state['model_state'].values() if v.is_floating_point())
        del state
        ds=row['dataset']; first={g:0 for g in GROUPS}
        for g,conditions in GROUPS.items():
            for c in conditions:
                for m in METRICS:
                    v=row['table'][g][c][m];assert math.isfinite(v)
                    first[g]+=v<min(b['table'][g][c][m] for n,b in base[ds].items() if n!='ALFNet')
                for f,v in base[ds]['ALFNet']['table'][g][c].items():
                    if f.endswith('_count'):assert row['table'][g][c][f]==v
        assert sum(first.values())==row['first_cells']
        audited[key]=dict(dataset=ds,model='ALFNet',seed=32,table=row['table'],selected_run=key,
            config=row['config'],config_sha256=config_sha,training_config_sha256=report['config_sha256'],plan_sha256=report['plan_sha256'],
            checkpoint=row['checkpoint'],checkpoint_sha256=row['checkpoint_sha256'],
            report=row['report'],report_sha256=study.sha(rp),best_epoch=report['best_epoch'],epochs_run=report['epochs_run'])
        candidates.append(dict(run=key,dataset=ds,selected=SELECTED[ds]==key,**params(row),
            **{m:row['table']['main']['LifeAvg'][m] for m in METRICS},
            main_first=first['main'],phase_first=first['phase'],first_cells=sum(first.values()),
            best_epoch_zero_based=report['best_epoch'],epochs_run=report['epochs_run']))
    for ds,key in SELECTED.items():records[ds]['ALFNet']=audited[key]
    summary={}
    for ds,models in records.items():
        a=models['ALFNet']['table'];old=base[ds]['ALFNet']['table']
        first={g:sum(rank(models,'ALFNet',g,c,m)==1 for c in cs for m in METRICS) for g,cs in GROUPS.items()}
        oldfirst=sum(rank(base[ds],'ALFNet',g,c,m)==1 for g,cs in GROUPS.items() for c in cs for m in METRICS)
        summary[ds]=dict(selected=SELECTED[ds],reason=REASONS[ds],first=first,old_first=oldfirst,
            lifeavg={m:a['main']['LifeAvg'][m] for m in METRICS},
            delta_percent={m:100*(a['main']['LifeAvg'][m]/old['main']['LifeAvg'][m]-1) for m in METRICS},
            lifeavg_ranks={m:rank(models,'ALFNet','main','LifeAvg',m) for m in METRICS})
        assert all(v==1 for v in summary[ds]['lifeavg_ranks'].values())
    OUT.mkdir(parents=True,exist_ok=True)
    notes=['五数据集、seed32、validation；ALFNet＋十个基线。第一名加粗，第二个不同数值下划线，按原精度排名。',
        '55个ALFNet候选（31复用＋24新增）；按33格覆盖和LifeAvg三指标综合取舍，非预注册单一评分规则，不设旧硬退化门槛。',
        '全部候选保留；33格有重叠，不是33次独立重复。单种子结果不作显著性推断。',
        '本次仅更新validation选型，既有test表和发布权重保持历史身份；不得混作新配置test结果。',
        '十个基线数值及来源保持不变。MoGERNN既有结果patience=3/start20，非10；SLPF单阶段、STBP离线结构编码，其他适配差异见../BASELINE_FAIRNESS_20260928.md。']
    wb=Workbook();wb.active.title='说明'
    for note in notes:wb.active.append([note])
    md=['# 五数据集最终选型 validation 主表与阶段表','',*notes,'']
    cells=[]; checks=[]
    for ds,models in records.items():
        for g,conditions in GROUPS.items():
            sheet=wb.create_sheet(ds+'_'+g);keys=[(c,m) for c in conditions for m in METRICS]
            header=['模型']+[f'{c}/{m}' for c,m in keys];sheet.append(header)
            md += [f'## {ds} '+('主表' if g=='main' else '阶段表'),'', '| '+' | '.join(header)+' |','|'+'|'.join(['---']+['---:']*len(keys))+'|']
            for ri,name in enumerate(['ALFNet']+[n for n in models if n!='ALFNet'],2):
                sheet.cell(ri,1,name);shown=[]
                for ci,(c,m) in enumerate(keys,2):
                    value=models[name]['table'][g][c][m];rk=rank(models,name,g,c,m)
                    cell=sheet.cell(ri,ci,value);cell.number_format='0.000';cell.font=Font(bold=rk==1,underline='single' if rk==2 else None)
                    txt=f'{value:.3f}';shown.append(f'**{txt}**' if rk==1 else f'<u>{txt}</u>' if rk==2 else txt)
                    cells.append([ds,name,g,c,m,value,rk]);checks.append((sheet.title,ri,ci,value,rk))
                md.append('| '+name+' | '+' | '.join(shown)+' |')
            sheet.freeze_panes='B2';sheet.column_dimensions['A'].width=18;md.append('')
    selected=[next(c for c in candidates if c['run']==SELECTED[ds]) for ds in SELECTED]
    hp=['# ALFNet 五数据集超参数表','',
        '下表对应本次validation选型的实际配置。容量是每个空间分区的token容量，分区数固定32；学习率倍率同时作用于图、时间、门控三个参数组。Dropout指ALFTemporal的crossformer_dropout；权重指AdamW weight_decay，不是损失项权重。','',
        '| 数据集 | Patch容量 | LR倍率 | 图分支LR | 时间分支LR | 门控LR | Dropout | Weight decay | 最佳epoch（从1计） | 总epoch |',
        '|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|']
    hs=wb.create_sheet('ALFNet超参数');hkeys=['dataset','capacity','lr_multiplier','graph_lr','temporal_lr','gate_lr','dropout','weight_decay','batch_size','max_epoch','patience','early_stopping_start_epoch','loss','best_epoch_zero_based','epochs_run','run']
    hs.append(hkeys)
    for c in selected:
        hs.append([c[k] for k in hkeys]);hp.append(f"| {c['dataset']} | {c['capacity']} | ×{c['lr_multiplier']:g} | {c['graph_lr']:g} | {c['temporal_lr']:g} | {c['gate_lr']:g} | {c['dropout']:g} | {c['weight_decay']:g} | {c['best_epoch_zero_based']+1} | {c['epochs_run']} |")
    hp += ['', '共同配置：seed32、batch64、AdamW、max_epoch100、patience10、early_stopping_start_epoch40、Focal(mae_focal_raw02)、12步输入/12步预测、teacher forcing线性衰减40轮；参数量1,037,529。', '',
        '每个数据集复用历史搜索并在其已选锚点补点，并非全部参数组合的网格搜索。比较单变量效应时须固定其他参数；完整55项实值及结果见all_candidates.csv和工作簿“全部超参数实验”。']
    cs=wb.create_sheet('全部超参数实验');ckeys=list(candidates[0]);cs.append(ckeys)
    for c in candidates:cs.append([c[k] for k in ckeys])
    (OUT/'final_tables.md').write_text('\n'.join(md)+'\n');(OUT/'ALFNET_HYPERPARAMETERS.md').write_text('\n'.join(hp)+'\n')
    wb.save(OUT/'final_tables.xlsx')
    loaded=load_workbook(OUT/'final_tables.xlsx')
    for name,ri,ci,value,rk in checks:
        cell=loaded[name].cell(ri,ci)
        assert abs(cell.value-value)<1e-10 and cell.font.bold==(rk==1) and cell.font.underline==('single' if rk==2 else None)
    for ri,c in enumerate(selected,2):
        for ci,k in enumerate(hkeys,1):assert loaded['ALFNet超参数'].cell(ri,ci).value==c[k]
    assert len(cells)==1815
    for filename,headers,values in [('cells.csv',['dataset','model','section','condition','metric','value','rank'],cells),('all_candidates.csv',ckeys,[[c[k] for k in ckeys] for c in candidates]),('alfnet_hyperparameters.csv',hkeys,[[c[k] for k in hkeys] for c in selected])]:
        with (OUT/filename).open('w',newline='') as f:w=csv.writer(f);w.writerow(headers);w.writerows(values)
    study.write(OUT/'results_and_provenance.json',dict(split='validation',seed=32,plan_sha256=p['plan_sha256'],records=records,summary=summary,selected=SELECTED,source=str(source.relative_to(ROOT)),source_sha256=study.sha(source),baseline_source=str(study.BASE_TABLE.relative_to(ROOT)),baseline_source_sha256=study.sha(study.BASE_TABLE),test_accessed=False))
    study.write(OUT/'verification.json',dict(passed=True,audited_candidates=55,new_runs_complete=24,checked_candidate_cells=55*33,excel_cells_verified=len(cells),selected_hparams_verified=True,checkpoint_hashes_and_finiteness=True,epoch_histories_contiguous=True,baseline_records_unchanged=all(records[d][n]==v for d,ms in base.items() for n,v in ms.items() if n!='ALFNet'),test_accessed=False))
    analysis=['# 超参数完成与最终选型分析','',*notes,'',
        '24/24新增训练exit0；55项候选的报告、配置、checkpoint哈希、最佳epoch、连续训练历史、有限权重及无跳过batch均核验通过。断电续训历史连续。','',
        '| 数据集 | LifeAvg MAE | RMSE | MAPE (%) | 主表第一/15 | 阶段第一/18 | 第一格：旧→新 |',
        '|---|---:|---:|---:|---:|---:|---:|']
    for ds,s in summary.items():
        analysis.append('| '+ds+' | '+' | '.join(f"{s['lifeavg'][m]:.4f}" for m in METRICS)+f" | {s['first']['main']} | {s['first']['phase']} | {s['old_first']}→{sum(s['first'].values())} |")
    analysis += ['', '五数据集LifeAvg三指标全部保持第一（15/15）；总第一格53→69/165。各数据集仍未达到80%第一格。平均收益与逐条件覆盖是不同目标；不能把所有候选的最小单元格拼成一个模型。','', '## 配置取舍（相对上一版validation表）','']
    for ds,s in summary.items():
        analysis += [f"- **{ds}**：{s['reason']} LifeAvg变化（负为改善）："+' / '.join(f"{m} {s['delta_percent'][m]:+.2f}%" for m in METRICS)+'。']
    analysis += ['', '## 搜索结论与备选','',
        'METR-LA：学习率×2是候选LifeAvg三指标的唯一Pareto点，且25/33格第一。',
        'PEMS-BAY：容量16/学习率×1是LifeAvg三指标唯一Pareto点；×2提升覆盖但不称平均误差更优。此处主动接受取舍，非硬门槛筛选。',
        'PEMS03：容量20较有利，学习率减半主要改善RMSE/MAPE；第一格数量仍受强基线限制。',
        'PEMS04：新dropout、weight decay及容量补点未超过原配置的覆盖与MAE/RMSE组合，因此不替换。',
        'PEMS08：dropout0.05/0.2和weight_decay1e-5构成LifeAvg Pareto集合；0.05更适合当前覆盖目标。旧LR×2虽11格第一，相对本次所选MAE/RMSE/MAPE增加约5.15%/5.31%/26.70%，不保留为首选。',
        '', '搜索为单种子、分阶段单变量补点，不构成全局最优证明；不跨speed/flow平均绝对误差。新配置的test表现未在本轮评估，旧test结论不能自动迁移。', '',
        '输出：final_tables.xlsx同时包含十张主/阶段表、五数据集超参数表和55项完整搜索；Markdown及CSV可独立使用。所有旧表、原始权重和配置均保留。']
    (OUT/'RESULTS_ANALYSIS.md').write_text('\n'.join(analysis)+'\n')
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
