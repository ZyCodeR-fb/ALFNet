"""Explicitly authorized test of 48 fixed ablations, preceded by validation replay."""
import argparse,copy,sys
from pathlib import Path
from contextlib import ExitStack
from unittest.mock import patch
import yaml
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.run_metrla_hparams import read,write,sha,digest
from scripts import test_all_final_20261001 as evaluator
from scripts import evaluate_speed_support_20261006 as queue_base
from scripts import run_final_ablations_20261006 as base
from scripts import run_mechanism_ablations_20261008 as mechanism
from scripts import run_reconstruction_gradient_20261009 as gradient
NAME='all_ablations_test_20261010';OUT=ROOT/'paper'/NAME;PLAN=ROOT/'configs/paper'/NAME/'plan.json'
SOURCE=ROOT/'paper/ablations_complete_20261010/results_and_audit.json'
def output(p):return ROOT/'artifacts/evaluation'/NAME/p['plan_sha256']
def freeze():
    import torch
    assert not PLAN.exists();base.check();mechanism.check();gradient.check()
    src=read(SOURCE);assert src['all_runs_complete'] and len(src['rows'])==52
    entries={};files=[SOURCE,OUT/'TEST_PLAN.md'];full={}
    old=ROOT/'paper/all_final_test_20261001/all_results_and_provenance.json';flow=ROOT/'paper/flow_lr_cross_test_20261004/test_results.json'
    manifest=ROOT/'paper/final_scheme_20261006/selected_configuration_manifest.json';selected=read(manifest)['entries'];files.extend([old,flow,manifest])
    for ds,v in selected.items():
        t=(read(old)['rows'] if ds in ('METRLA','PEMSBAY') else read(flow)['rows'])[v['key']]
        assert t['passed'] and t['checkpoint_sha256']==sha(ROOT/src['rows'][ds+'_Full']['checkpoint'])
        full[ds+'_Full']=dict(t,label='Full')
    for key,row in src['rows'].items():
        if row['label']=='Full':continue
        r=copy.deepcopy(row);cp=torch.load(ROOT/r['checkpoint'],map_location='cpu',weights_only=False)
        assert cp['seed']==32 and cp['config_sha256']==sha(ROOT/r['config'])
        assert sha(ROOT/r['checkpoint'])==src['audits'][key]['checkpoint_sha256']
        cfg=yaml.safe_load((ROOT/r['config']).read_text());kind=cfg['family']['recipe']['study_ablation']
        r.update(model='ALFNet',epoch=cp['epoch'],training_config_sha256=cp['config_sha256'],training_plan_sha256=cp['plan_sha256'],checkpoint_sha256=sha(ROOT/r['checkpoint']),ablation=kind)
        entries[key]=r
        files.extend(ROOT/r[f] for f in ('config','checkpoint','report'))
        files.extend(ROOT/cfg['family']['experts'][e]['config_path'] for e in ('graph','patch'))
        files.extend(ROOT/cfg['dataset'][f] for f in ('data_path','unknown_nodes_path'))
        files.append(ROOT/f"configs/paper/v4_data_manifests/{r['dataset']}_run_manifest.json")
    for folder in ('model','layers','experiments','lib','data_provider','utils','scripts'):files.extend((ROOT/folder).rglob('*.py'))
    files.append(ROOT/'args.py')
    p=dict(kind=NAME,status='authorized_test',test_authorized=True,authorization='2026-10-10 user explicitly requests test all previously completed ablations and evaluate four-dataset LifeAvg criterion',
        entries=entries,reused_full=full,max_concurrent=3,reporting='All48 frozen ablations retained, validation replay before test, no outcome-based selection or model change',files={str(f.relative_to(ROOT)):sha(f) for f in sorted(set(files))})
    p['plan_sha256']=digest(p);write(PLAN,p);print(p['plan_sha256'])
def check():
    p=read(PLAN);assert digest(p)==p['plan_sha256'] and p['test_authorized'] and len(p['entries'])==48 and len(p['reused_full'])==4
    for f,h in p['files'].items():assert sha(ROOT/f)==h,f
    return p
def worker(key):
    import torch
    from experiments import v4_runner
    from scripts import train_v4
    p=check();r=p['entries'][key];cp=torch.load(ROOT/r['checkpoint'],map_location='cpu',weights_only=False)
    with ExitStack() as stack:
        if r['ablation'] in gradient.LABELS:
            stack.enter_context(gradient.scope());stack.enter_context(mechanism.scope())
        elif r['ablation'] in mechanism.LABELS:stack.enter_context(mechanism.scope())
        stack.enter_context(base.context());build=v4_runner.build_v4_session
        def verified(*args,**kwargs):
            s=build(*args,**kwargs);assert train_v4.model_method_contract(s)==cp['model_method_contract'];return s
        stack.enter_context(patch.object(v4_runner,'build_v4_session',verified))
        stack.enter_context(patch.object(evaluator,'check',check));stack.enter_context(patch.object(evaluator,'output',output));evaluator.worker(key)
    path=output(p)/(key+'.json');result=read(path);result.update(label=r['label'],ablation=r['ablation']);write(path,result)
def collect():
    from scripts.package_supplements_20261008 import book_from_cells,csvwrite
    from openpyxl import load_workbook
    p=check();root=output(p);rows=copy.deepcopy(p['reused_full']);deltas={};replays={}
    for k,e in p['entries'].items():
        r=read(root/(k+'.json'));v=read(root/'validation'/(k+'.json'))
        assert r['passed'] and v['passed'] and r['checkpoint_sha256']==e['checkpoint_sha256']
        rows[k]=r;replays[k]=dict(passed=True,max_abs_difference=v['max_abs_difference'])
    cells=[];metrics=('MAE','RMSE','MAPE');datasets=('METRLA','PEMSBAY','PEMS04','PEMS08')
    for k,r in rows.items():
        ds=r['dataset'];ref=rows[ds+'_Full']['table']
        for g in ('main','phase'):
            for c in ref[g]:
                for f,v in ref[g][c].items():
                    if f.endswith('_count'):assert r['table'][g][c][f]==v
                for m in metrics:cells.append(dict(dataset=ds,run=k,label=r['label'],section=g,condition=c,metric=m,value=r['table'][g][c][m]))
        if r['label']!='Full':deltas.setdefault(r['label'],{})[ds]={m:100*(r['table']['main']['LifeAvg'][m]/ref['main']['LifeAvg'][m]-1) for m in metrics}
    assert len(deltas)==12 and all(len(v)==4 for v in deltas.values())
    qualifying=[l for l,s in deltas.items() if all(v>0 for d in s.values() for v in d.values())]
    n=book_from_cells(OUT/'all_ablations_test.xlsx',cells,['48 fixed ablations+4Full,seed32 test; all results retained; matched Full checkpoint.', 'First bold/second underline; lower error is better.'],['dataset','section'],['label']);assert n==1716
    wb=load_workbook(OUT/'all_ablations_test.xlsx');ws=wb.create_sheet('LifeAvg变化及判定');ws.append(['消融']+[d+'/'+m+'%' for d in datasets for m in metrics]+['变差指标/12','四数据集全部符合'])
    for label,s in deltas.items():ws.append([label]+[s[d][m] for d in datasets for m in metrics]+[sum(v>0 for d in s.values() for v in d.values()),label in qualifying])
    wb.save(OUT/'all_ablations_test.xlsx');csvwrite(OUT/'all_ablations_test.csv',cells,list(cells[0]))
    write(OUT/'results_and_audit.json',dict(rows=rows,validation_replays=replays,lifeavg_delta_percent=deltas,qualifying=qualifying,qualified_count=len(qualifying),plan_sha256=p['plan_sha256'],metric_cells=n,test_accessed=True))
    lines=['# 全部消融test汇总','',f'48项test完成，4项Full复用；1716指标格。符合四数据集三项LifeAvg均严格变差：{len(qualifying)}种。',
        '全部结果保留；validation和test分开解释，本轮不按test更新模型或剔除消融。','', '| 消融 | METRLA变化% | PEMSBAY | PEMS04 | PEMS08 | 全12项变差 |','|---|---|---|---|---|---|']
    for label,s in deltas.items():lines.append('| '+label+' | '+' | '.join('/'.join(f'{s[d][m]:+.2f}' for m in metrics) for d in datasets)+f' | {label in qualifying} |')
    (OUT/'ANALYSIS.md').write_text('\n'.join(lines)+'\n');write(root/'COMPLETE.json',dict(passed=True,new_tests=48,reused_full=4,metric_cells=n,qualified_count=len(qualifying)))
def queue():
    p=check()
    with ExitStack() as stack:
        for n,v in [('check',check),('output',output),('collect',collect),('KEYS',tuple(p['entries'])),('__file__',__file__)]:stack.enter_context(patch.object(queue_base,n,v))
        queue_base.queue()
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['freeze','check','worker','queue','collect']);p.add_argument('--key');a=p.parse_args()
    if a.mode=='worker':worker(a.key)
    elif a.mode=='check':print(check()['plan_sha256'])
    else:globals()[a.mode]()
