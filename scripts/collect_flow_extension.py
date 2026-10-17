"""Offline final validation tables; no model execution or test access."""
import copy
import argparse
import csv
import math
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.run_flow_extension import check, output, read, write, sha

OUT = ROOT / 'paper/validation_seed32_flow_20260927'
SPEED = ROOT / 'paper/validation_seed32_hparam_patch16_20260918/results_and_provenance.json'
METRICS = ('MAE', 'RMSE', 'MAPE')
GROUPS = {'main': ('Unseen', 'Steady', 'Join', 'Fault', 'LifeAvg'),
          'phase': ('Cold', 'Warm', 'Ready', 'Onset', 'Decay', 'Outage')}


def main(revised=False):
    global OUT
    if revised:
        OUT = ROOT / 'paper/validation_seed32_flow_selected_20260927'
    p = check()
    root = output(p)
    results = read(ROOT / 'paper/flow_extension_20260926/results.json')
    assert results['plan_sha256'] == p['plan_sha256'] and results['split'] == 'validation'
    assert read(root / 'COMPLETE.json')['passed']
    queue = read(root / 'queue_status.json')
    assert len(queue['jobs']) == 15 and all(j['status'] == 'complete' and j['exit_code'] == 0 for j in queue['jobs'])
    records = copy.deepcopy(read(SPEED)['records'])
    provenance, audits, summaries = {}, [], []
    for key, entry in p['configs'].items():
        done = read(root / (key + '_complete.json'))
        assert done['passed'] and done['plan_sha256'] == p['plan_sha256']
        report_path = Path(done['report_path'])
        report = read(report_path)
        assert report['config_sha256'] == entry['sha256'] and report['plan_sha256'] == p['plan_sha256']
        assert all(e[phase]['skipped'] == 0 for e in report['history'] for phase in ('train', 'validation'))
        scores = [(e['validation']['checkpoint_life_avg_mae'], e['epoch']) for e in report['history']]
        assert min(scores)[1] == report['best_epoch']
        table = report['development_validation_table']
        assert table == results['tables'][key]
        baseline = results['tables'][entry['dataset'] + '_MoGERNN']
        wins = {}
        for group, conditions in GROUPS.items():
            wins[group] = 0
            for condition in conditions:
                for m in METRICS:
                    value = table[group][condition][m]
                    assert math.isfinite(value)
                    wins[group] += value < baseline[group][condition][m]-1e-6
                for field, value in baseline[group][condition].items():
                    if field.endswith('_count'):
                        assert table[group][condition][field] == value
        assert wins['main'] == results['wins_vs_mogernn'][key]['main']
        assert wins['phase'] == results['wins_vs_mogernn'][key]['phase']
        checkpoint = report_path.with_name('best.pt')
        assert checkpoint.exists() and report_path.with_name('COMPLETE').exists()
        provenance[key] = dict(report=str(report_path.relative_to(ROOT)), report_sha256=sha(report_path),
                               checkpoint=str(checkpoint.relative_to(ROOT)), checkpoint_sha256=sha(checkpoint),
                               config=entry['path'], config_sha256=entry['sha256'],
                               best_epoch=report['best_epoch'], epochs_run=report['epochs_run'])
        summaries.append([key, entry['capacity'], entry['lr_multiplier'], report['best_epoch'], report['epochs_run'],
                          *[table['main']['LifeAvg'][m] for m in METRICS], wins['main'], wins['phase']])
        audits.append(dict(run=key, cells_checked=33, best_epoch_verified=True, no_skipped_batches=True))
    for ds, decision in results['decisions'].items():
        keys = [k for k, v in p['configs'].items() if v['dataset'] == ds]
        best = min(keys, key=lambda k: results['tables'][k]['main']['LifeAvg']['MAE'])
        assert best == decision['minimum_mae_candidate']
        t, b = results['tables'][best], results['tables'][ds + '_default']
        delta = {m: t['main']['LifeAvg'][m]/b['main']['LifeAvg'][m]-1 for m in METRICS}
        early = sum(t['phase'][c]['MAE'] for c in ('Cold', 'Warm', 'Ready'))/sum(b['phase'][c]['MAE'] for c in ('Cold', 'Warm', 'Ready'))-1
        unseen = t['main']['Unseen']['MAE']/b['main']['Unseen']['MAE']-1
        passed = delta['MAE'] <= -.005 and delta['RMSE'] <= .005 and delta['MAPE'] <= .005 and early <= 0 and unseen <= .01
        retained = best if passed else ds + '_default'
        assert retained == decision['retained'] and passed == decision['passes_screen']
        if revised:
            retained = best
            decision['previous_retained'] = decision['retained']
            decision['retained'] = retained
            decision['selection_policy'] = '2026-09-27 user cancels regression gates; validation LifeAvg MAE with full three-metric and win-count review'
        reuse = read(root / (ds + '_reuse.json'))
        assert reuse['passed'] and reuse['max_validation_delta'] == 0 and reuse['table'] == results['tables'][ds + '_MoGERNN']
        records[ds] = {'ALFNet': dict(table=results['tables'][retained], selected_run=retained, **provenance[retained]),
                       'MoGERNN': dict(table=reuse['table'], source=reuse['source'], reuse_verification=str((root / (ds + '_reuse.json')).relative_to(ROOT)))}
    OUT.mkdir(parents=True, exist_ok=True)
    notes = ['五数据集 seed32 validation 主表与阶段表。speed 两数据集复用原十模型记录；flow 三数据集仅包含本轮 ALFNet 和 MoGERNN。',
             '本表不是 test 表，不能与正式 test 数值混合比较。flow ALFNet 按冻结门槛保留配置，其他候选完整保存在候选汇总中。',
             '第一名加粗，第二个不同数值下划线；按未舍入值排序。flow 的排名仅针对两个模型。',
             'MAE/RMSE为各数据集原单位，MAPE为%；不跨 speed/flow 平均绝对误差。单种子结果不代表统计显著性。',
             'SLPF沿用单阶段预测适配，STBP沿用结构编码/离线固定图适配；本轮没有重新运行或恢复原论文完整训练方法。']
    if revised:
        notes[1] = '本表为validation。2026-09-27用户取消旧硬性退化门槛；按LifeAvg MAE结合三指标取舍保留配置，三者均对MoGERNN达到33/33胜。旧表和原计划保留。'
    wb = Workbook()
    wb.active.title = '说明'
    for note in notes:
        wb.active.append([note])
    md = ['# 最终保留配置：五数据集 validation 主表与阶段表', '', *notes, '']
    checks, cells = [], []
    for ds, items in records.items():
        names = ['ALFNet'] + [n for n in items if n != 'ALFNet']
        for group, conditions in GROUPS.items():
            title = ds + ('主表' if group == 'main' else '阶段表')
            sheet = wb.create_sheet(title)
            keys = [(c, m) for c in conditions for m in METRICS]
            header = ['模型'] + [c + '/' + m for c, m in keys]
            sheet.append(header)
            md += ['## ' + title, '', '| ' + ' | '.join(header) + ' |', '|' + '|'.join(['---'] + ['---:']*len(keys)) + '|']
            for ri, name in enumerate(names, 2):
                label = {'STFiT': 'ST-FiT', 'SLPF': 'SLPF（单阶段）'}.get(name, name)
                sheet.cell(ri, 1, label)
                shown = []
                for ci, (c, m) in enumerate(keys, 2):
                    value = items[name]['table'][group][c][m]
                    rank = sorted(set(items[n]['table'][group][c][m] for n in names)).index(value)+1
                    cell = sheet.cell(ri, ci, value)
                    cell.number_format = '0.000'
                    cell.font = Font(bold=rank == 1, underline='single' if rank == 2 else None)
                    text = f'{value:.3f}'
                    shown.append(f'**{text}**' if rank == 1 else f'<u>{text}</u>' if rank == 2 else text)
                    checks.append((title, ri, ci, value, rank))
                    cells.append([ds, name, group, c, m, value, rank])
                md.append('| ' + label + ' | ' + ' | '.join(shown) + ' |')
            sheet.freeze_panes = 'B2'
            sheet.column_dimensions['A'].width = 25
            md.append('')
    (OUT / 'final_tables.md').write_text('\n'.join(md))
    wb.save(OUT / 'final_tables.xlsx')
    loaded = load_workbook(OUT / 'final_tables.xlsx')
    for sheet, row, col, value, rank in checks:
        c = loaded[sheet].cell(row, col)
        assert abs(c.value-value) <= 1e-12*max(1, abs(value))
        assert c.font.bold == (rank == 1) and (c.font.underline == 'single') == (rank == 2)
    for filename, headers, rows in [('cells.csv', ['dataset', 'model', 'section', 'condition', 'metric', 'value', 'rank'], cells),
                                    ('all_candidates.csv', ['run', 'capacity', 'lr_multiplier', 'best_epoch', 'epochs_run', 'MAE', 'RMSE', 'MAPE', 'main_wins', 'phase_wins'], summaries)]:
        with (OUT / filename).open('w', newline='') as f:
            w = csv.writer(f); w.writerow(headers); w.writerows(rows)
    write(OUT / 'results_and_provenance.json', dict(split='validation', plan_sha256=p['plan_sha256'],
          records=records, candidates=provenance, decisions=results['decisions'],
          speed_source=str(SPEED.relative_to(ROOT)), speed_source_sha256=sha(SPEED), test_accessed=False))
    write(OUT / 'verification.json', dict(passed=True, plan_sha256=p['plan_sha256'], runs=audits,
          raw_candidate_cells=495, published_cells=len(checks), excel_numeric_and_rank_verified=True,
          speed_records_unchanged=all(records[k] == v for k, v in read(SPEED)['records'].items()),
          selection_recomputed=True, test_accessed=False))
    print(f'Verified 15 runs, 495 candidate cells, {len(checks)} published cells; no test access.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--revised-selection', action='store_true')
    main(parser.parse_args().revised_selection)
