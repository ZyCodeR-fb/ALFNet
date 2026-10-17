"""Versioned combination results; conservative automatic updates, tradeoffs visible."""
import copy
import csv
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from scripts.run_metrla_hparams import read,write,sha
from scripts.collect_hparams_final import GROUPS,METRICS,params,rank


def unambiguous(candidate,current):
    a=[candidate['table']['main']['LifeAvg'][m] for m in METRICS]
    b=[current['table']['main']['LifeAvg'][m] for m in METRICS]
    return (candidate['first_cells']>=current['first_cells'] and all(x<=y for x,y in zip(a,b))
            and (candidate['first_cells']>current['first_cells'] or any(x<y for x,y in zip(a,b))))


def publish():
    from scripts import run_hparam_combinations as run
    p=run.check();out=ROOT/'paper'/run.NAME
    data=read(out/'results.json');rows=data['rows'];assert len(rows)==67
    prior=read(run.PRIOR/'results_and_provenance.json');records=copy.deepcopy(prior['records'])
    selected={ds:r['ALFNet']['selected_run'] for ds,r in records.items()};decisions={}
    for ds,key in list(selected.items()):
        current=rows[key]
        new=[k for k,e in p['configs'].items() if e['dataset']==ds]
        eligible=[k for k in new if unambiguous(rows[k],current)]
        best=min(eligible,key=lambda k:(-rows[k]['first_cells'],rows[k]['table']['main']['LifeAvg']['MAE'])) if eligible else key
        decisions[ds]=dict(previous=key,automatic_selected=best,unambiguous=eligible,
            tradeoffs_for_review=[k for k in new if k not in eligible],automatic_rule_is_not_rejection_gate=True)
        selected[ds]=best
        if best!=key:
            r=rows[best];report=read(ROOT/r['report'])
            records[ds]['ALFNet']=dict(dataset=ds,model='ALFNet',seed=32,selected_run=best,
                table=r['table'],config=r['config'],config_sha256=report['config_sha256'],
                checkpoint=r['checkpoint'],checkpoint_sha256=sha(ROOT/r['checkpoint']),
                report=r['report'],report_sha256=sha(ROOT/r['report']),plan_sha256=report['plan_sha256'],
                best_epoch=report['best_epoch'],epochs_run=report['epochs_run'])
    wb=load_workbook(run.PRIOR/'final_tables.xlsx');wb['说明'].append(['本轮12个组合结果；自动更新只采用无取舍改善，其他候选保留待综合审阅。validation，不是test。'])
    md=['# 最后一轮组合：版本化validation主表/阶段表','',
        '仅无取舍改善自动更新；有取舍候选并未被淘汰，见selection_review.json和全部67项结果。历史单变量sheet保持原参考点。','']
    cells=[];checks=[]
    for ds,models in records.items():
        for g,conditions in GROUPS.items():
            ws=wb[ds+'_'+g];names=[ws.cell(i,1).value for i in range(2,13)]
            keys=[(c,m) for c in conditions for m in METRICS]
            md += [f'## {ds} {g}','', '| 模型 | '+' | '.join(c+'/'+m for c,m in keys)+' |','|'+'|'.join(['---']+['---:']*len(keys))+'|']
            for ri,name in enumerate(names,2):
                text=[]
                for ci,(c,m) in enumerate(keys,2):
                    v=models[name]['table'][g][c][m];rk=rank(models,name,g,c,m)
                    cell=ws.cell(ri,ci,v);cell.font=Font(bold=rk==1,underline='single' if rk==2 else None)
                    t=f'{v:.3f}';text.append(f'**{t}**' if rk==1 else f'<u>{t}</u>' if rk==2 else t)
                    cells.append([ds,name,g,c,m,v,rk]);checks.append((ws.title,ri,ci,v,rk))
                md.append('| '+name+' | '+' | '.join(text)+' |')
            md.append('')
    hp=wb['ALFNet超参数'];headers=[c.value for c in hp[1]];hprows=[]
    for ri,ds in enumerate(prior['records'],2):
        key=selected[ds];r=rows[key];report=read(ROOT/r['report']);v=dict(dataset=ds,run=key,**params(r),best_epoch_zero_based=report['best_epoch'],epochs_run=report['epochs_run'])
        for ci,h in enumerate(headers,1):hp.cell(ri,ci,v[h])
        hprows.append(v)
    for ds in records:
        ws=wb.create_sheet('最后组合_'+ds)
        columns=['run','容量','LR倍率','Dropout','WD','第一格','当前自动保留']+[g+'/'+c+'/'+m for g,cs in GROUPS.items() for c in cs for m in METRICS]
        ws.append(columns)
        keys=[prior['records'][ds]['ALFNet']['selected_run']]+[k for k,e in p['configs'].items() if e['dataset']==ds]
        for key in keys:
            r=rows[key];v=params(r)
            ws.append([key,*[v[k] for k in ('capacity','lr_multiplier','dropout','weight_decay')],r['first_cells'],key==selected[ds]]+[r['table'][g][c][m] for g,cs in GROUPS.items() for c in cs for m in METRICS])
        ws.freeze_panes='H2';ws.column_dimensions['A'].width=34
    wb.save(out/'final_tables.xlsx')
    loaded=load_workbook(out/'final_tables.xlsx')
    for name,ri,ci,v,rk in checks:
        c=loaded[name].cell(ri,ci);assert abs(c.value-v)<1e-10 and c.font.bold==(rk==1) and c.font.underline==('single' if rk==2 else None)
    (out/'final_tables.md').write_text('\n'.join(md)+'\n')
    with (out/'final_table_cells.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(['dataset','model','section','condition','metric','value','rank']);w.writerows(cells)
    write(out/'selection_review.json',dict(decisions=decisions,selected=selected,review_tradeoffs=True,test_accessed=False))
    write(out/'results_and_provenance.json',dict(records=records,selected=selected,split='validation',plan_sha256=p['plan_sha256'],test_accessed=False))
    write(out/'alfnet_hyperparameters.json',hprows)
    write(out/'publication_verification.json',dict(passed=True,excel_cells=len(checks),old_baselines_unchanged=all(records[d][n]==v for d,ms in prior['records'].items() for n,v in ms.items() if n!='ALFNet'),test_accessed=False))
    write(run.output(p)/'TABLES_COMPLETE.json',dict(passed=True,output=str(out.relative_to(ROOT))))


if __name__=='__main__':publish()
