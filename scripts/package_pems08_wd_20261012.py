"""Versioned selected tables; preserve previous evaluations and ablation anchors."""
import copy,csv,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.test_pems08_wd_20261012 import check,output,KEY
from scripts.run_metrla_hparams import read,write,sha
from scripts.audit_full_vs_five_ablations_20261012 import compare,BASE,KINDS
from scripts import plot_advisor_experiments_20261012 as plot
OUT=ROOT/'paper/pems08_wd_test_20261012'
def main():
    import numpy as np
    import matplotlib.pyplot as plt
    from matplotlib.colors import TwoSlopeNorm
    from openpyxl import Workbook
    from openpyxl.styles import Font
    p=check();r=read(output(p)/(KEY+'.json'));v=read(output(p)/'validation'/(KEY+'.json'))
    assert r['passed'] and v['passed'] and r['checkpoint_sha256']==p['entries'][KEY]['checkpoint_sha256']
    ab=read(ROOT/'paper/all_ablations_test_20261010/results_and_audit.json')['rows']
    old=read(ROOT/'paper/all_final_test_20261001/all_results_and_provenance.json')['rows']
    av=read(ROOT/'configs/paper/final_ablations_20261006/plan.json')['anchors']
    bv=read(ROOT/'paper/hparams_combinations_20260930/results_and_provenance.json')['records']
    test_bases={m:old['PEMS08_'+m]['table'] for m in BASE}
    abl={k:next(x['table'] for n,x in ab.items() if n.endswith('PEMS08_'+k)) for k in KINDS}
    ranking=compare(r['table'],test_bases,abl)
    previous=ab['PEMS08_Full']['table'];metrics=plot.MS
    changes={m:100*(r['table']['main']['LifeAvg'][m]/previous['main']['LifeAvg'][m]-1) for m in metrics}
    promoted=all(x<0 for x in changes.values()) and ranking['baseline_first']>=30
    decision=dict(promoted=promoted,plan_sha256=p['plan_sha256'],candidate=KEY,checkpoint_sha256=r['checkpoint_sha256'],test_ranking=ranking,delta_percent=changes,validation_replay_max_abs_difference=v['max_abs_difference'],previous_test_lifeavg=previous['main']['LifeAvg'],interpretation='Test-informed retention; comparisons against old WD=.0001 ablations are exploratory. New paired WD=.001 ablations required.')
    write(OUT/'decision.json',decision);assert promoted
    manifest=read(ROOT/'paper/final_scheme_20261006/selected_configuration_manifest.json')
    entry=manifest['entries']['PEMS08'];entry['previous_key']=entry['key'];entry['key']=KEY
    for k in ('config','checkpoint','checkpoint_sha256'):entry[k]=p['entries'][KEY][k]
    entry['config_sha256']=p['entries'][KEY]['training_config_sha256'];entry['hyperparameters']['weight_decay']=.001
    manifest.update(status='PEMS08 WD=.001 promoted after authorized test; other datasets unchanged',selection_evidence=str((OUT/'decision.json').relative_to(ROOT)))
    write(OUT/'selected_configuration_manifest.json',manifest)
    tables={s:{d:{m:(old[d+'_'+m]['table'] if s=='test' else bv[d][m]['table']) for m in BASE} for d in plot.DS} for s in ('validation','test')}
    for d in plot.DS:
        tables['test'][d]['ALFNet']=ab[d+'_Full']['table'];tables['validation'][d]['ALFNet']=av[d]['table']
    tables['test']['PEMS08']['ALFNet']=r['table'];tables['validation']['PEMS08']['ALFNet']=v['table']
    write(OUT/'selected_tables.json',tables)
    for split,datasets in tables.items():
        wb=Workbook();wb.remove(wb.active);cells=[]
        for d,models in datasets.items():
            for group in ('main','phase'):
                ws=wb.create_sheet(d+'_'+group);cols=[(c,m) for c in models['ALFNet'][group] for m in metrics]
                ws.append(['Model']+[c+'/'+m for c,m in cols])
                for name,t in models.items():
                    for c,m in cols:
                        for f,n in t[group][c].items():
                            if f.endswith('_count'):assert n==models['ALFNet'][group][c][f]
                    ws.append([name]+[t[group][c][m] for c,m in cols])
                    for j,(c,m) in enumerate(cols,2):
                        ranks=sorted(set(x[group][c][m] for x in models.values()));value=t[group][c][m]
                        ws.cell(ws.max_row,j).font=Font(bold=value==ranks[0],underline='single' if value==ranks[1] else None)
                        ws.cell(ws.max_row,j).number_format='0.000'
                        cells.append([d,name,group,c,m,value])
        wb.save(OUT/(split+'_main_phase.xlsx'))
        with (OUT/(split+'_main_phase.csv')).open('w') as f:
            w=csv.writer(f);w.writerow(['dataset','model','section','condition','metric','value']);w.writerows(cells)
    values={d:{n:t['main']['LifeAvg'] for n,t in ms.items()} for d,ms in tables['test'].items()}
    (OUT/'table1.md').write_text(plot.wide_table(values,list(BASE)+['ALFNet']))
    fig,axes=plt.subplots(2,2,figsize=(12,9),layout='constrained');heat=[];summary={}
    for ax,d,label in zip(axes.flat,plot.DS,plot.LABEL):
        ts=tables['test'][d];cols=[(g,c) for g in ('main','phase') for c in ts['ALFNet'][g]]
        a=np.array([[100*(1-ts['ALFNet'][g][c][m]/min(ts[b][g][c][m] for b in BASE)) for m in metrics] for g,c in cols])
        first=int((a>0).sum());summary[d]=dict(first=first,lifeavg_gain={m:100*(1-values[d]['ALFNet'][m]/min(values[d][b][m] for b in BASE)) for m in metrics})
        im=ax.imshow(a,cmap='RdBu',norm=TwoSlopeNorm(vmin=-40,vcenter=0,vmax=40),aspect='auto')
        for (i,j),x in np.ndenumerate(a):ax.text(j,i,f'{x:+.1f}',ha='center',va='center',fontsize=9,color='white' if abs(x)>28 else '#17252f')
        ax.set(xticks=range(3),xticklabels=metrics,yticks=range(11),yticklabels=[c for g,c in cols],title=f'{label} | {first}/33 best');ax.axhline(4.5,color='white',lw=2)
        heat.extend([d,g,c,m,float(a[i,j])] for i,(g,c) in enumerate(cols) for j,m in enumerate(metrics))
    fig.colorbar(im,ax=axes,shrink=.8,label='Error reduction vs. best baseline in each cell (%)')
    for ext in ('png','pdf','svg'):fig.savefig(OUT/('fig2_condition_stage.'+ext),dpi=200,bbox_inches='tight')
    plt.close(fig);write(OUT/'summary.json',summary)
    with (OUT/'fig2_values.csv').open('w') as f:
        w=csv.writer(f);w.writerow(['dataset','section','condition','metric','gain']);w.writerows(heat)
    lines=['# PEMS08 WD=0.001 确认与替换','','训练时间：北京时间2026-10-06 19:39:32—23:55:48，50 epochs，best epoch39。最终模型方案不变，仅WD由0.0001改为0.001。原先属固定锚点单因素实验，未自动选型、未进行test。','','完整validation回放差异为0。test三项LifeAvg均降低，第一名保持30/33，采用新配置。','','| 配置 | MAE | RMSE | MAPE |','|---|---:|---:|---:|', '| 原Full | '+' | '.join(f'{previous["main"]["LifeAvg"][m]:.6f}' for m in metrics)+' |','| WD=.001 | '+' | '.join(f'{ranking["lifeavg"][m]:.6f}' for m in metrics)+' |','','相对变化：'+str(changes),'','旧五项消融的15个LifeAvg值均高于新Full，但WD不同，不作为配对消融证据。五项新消融使用WD=.001从头训练，保留所有结果。旧配置、旧结果、旧计划保留。正文表2和机制图暂保留明确注明的历史固定锚点，不把旧权重结果改标成新权重。']
    (OUT/'ANALYSIS.md').write_text('\n'.join(lines)+'\n');write(output(p)/'COMPLETE.json',dict(passed=True,new_tests=1,promoted=True));print(decision);print(summary)
if __name__=='__main__':main()
