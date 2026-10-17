"""Audit completed training and package OFAT plus available paired ablations."""
import copy,csv,json,math,sys
from pathlib import Path
import torch
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.run_metrla_hparams import read,write,sha
from scripts.collect_hparams_final import params
OUT=ROOT/'paper/final_supplements_20261008'
METRICS=('MAE','RMSE','MAPE');DATASETS=('METRLA','PEMSBAY','PEMS04','PEMS08')
def audit_row(row,entry=None,plan=None):
    r=read(ROOT/row['report']);cp=ROOT/row['checkpoint'];cfg=ROOT/row['config']
    assert r['development_validation_table']==row['table']
    assert r['config_sha256']==sha(cfg)
    if entry:assert r['config_sha256']==entry['sha256'] and r['plan_sha256']==plan
    assert r['seed']==32 and all(e[s]['skipped']==0 for e in r['history'] for s in ('train','validation'))
    assert min((e['validation']['checkpoint_life_avg_mae'],e['epoch']) for e in r['history'])[1]==r['best_epoch']
    if row.get('checkpoint_sha256'):assert sha(cp)==row['checkpoint_sha256']
    state=torch.load(cp,map_location='cpu',weights_only=False)
    assert state['best_epoch']==r['best_epoch'] and state['config_sha256']==r['config_sha256'] and state['plan_sha256']==r['plan_sha256']
    assert all(torch.isfinite(v).all() for v in state['model_state'].values() if v.is_floating_point());del state
    assert all(math.isfinite(row['table'][g][c][m]) for g in ('main','phase') for c in row['table'][g] for m in METRICS)
    return dict(report_sha256=sha(ROOT/row['report']),checkpoint_sha256=sha(cp),best_epoch=r['best_epoch'],epochs_run=r['epochs_run'],passed=True)
def csvwrite(path,rows,fields):
    with path.open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def book_from_cells(path,cells,notes,groupfields,labelfields):
    wb=Workbook();ws=wb.active;ws.title='说明'
    for n in notes:ws.append([n])
    grouped={}
    for r in cells:grouped.setdefault(tuple(r[f] for f in groupfields),[]).append(r)
    verified=[]
    for i,(group,rs) in enumerate(grouped.items()):
        ws=wb.create_sheet('_'.join(str(g) for g in group)[:31])
        columns=list(dict.fromkeys((r['condition'],r['metric']) for r in rs));labels=list(dict.fromkeys(tuple(r[f] for f in labelfields) for r in rs))
        values={(tuple(r[f] for f in labelfields),r['condition'],r['metric']):float(r['value']) for r in rs}
        ws.append(list(labelfields)+[c+'/'+m for c,m in columns])
        for idx,label in enumerate(labels,2):
            ws.append(list(label)+[values[label,c,m] for c,m in columns])
            for j,(c,m) in enumerate(columns,len(labelfields)+1):
                val=values[label,c,m];ranks=sorted({values[l,c,m] for l in labels});rank=ranks.index(val)+1
                ws.cell(idx,j).font=Font(bold=rank==1,underline='single' if rank==2 else None);ws.cell(idx,j).number_format='0.0000'
                verified.append((ws.title,idx,j,val,rank))
        ws.freeze_panes=ws.cell(2,len(labelfields)+1).coordinate
    wb.save(path);check=load_workbook(path)
    for title,r,c,v,rank in verified:
        cell=check[title].cell(r,c);assert math.isclose(cell.value,v,rel_tol=1e-12,abs_tol=1e-12)
        assert cell.font.bold==(rank==1) and (cell.font.underline=='single')==(rank==2)
    return len(verified)
def main():
    torch.set_num_threads(4);OUT.mkdir(parents=True,exist_ok=True)
    from scripts import run_flow_ofat_20261006 as h,run_final_ablations_20261006 as a
    hp=h.check();ap=a.check();hr=h.output(hp);ar=a.output(ap)
    assert all(j['status']=='complete' and j['exit_code']==0 for j in read(hr/'queue_status.json')['jobs'])
    hrows=read(h.OUT/'results.json')['rows'];audits={'hyperparameters':{},'ablations':{}}
    for k,r in hrows.items():audits['hyperparameters'][k]=audit_row(r,hp['configs'].get(k),hp['plan_sha256'])
    # Preserve speed paper OFAT and every historical combination separately.
    old=list(csv.DictReader((ROOT/'paper/final_four_tables_20261001/03_hyperparameters_validation.csv').open()))
    fields=list(old[0]);primary=[r for r in old if r['dataset'] in ('METRLA','PEMSBAY') and r['use']=='正文单变量']
    history=[dict(r,use='历史配置/组合；不混入当前flow单因素曲线') for r in old if r not in primary]
    axes={'capacity':'capacity','lr_factor':'lr_multiplier','dropout':'dropout','weight_decay':'weight_decay'}
    for ds in ('PEMS04','PEMS08'):
        anchor=h.ANCHORS[ds]
        for axis,field in axes.items():
            keys=[anchor]+[k for k,r in hrows.items() if r['dataset']==ds and r.get('axis')==axis]
            # Reused doubled LR records carry the original metadata; set their new OFAT axis explicitly.
            if axis=='lr_factor' and h.DOUBLE[ds] not in keys:keys.append(h.DOUBLE[ds])
            keys=sorted(set(keys),key=lambda k:params(hrows[k])[field])
            assert len(keys)==3,(ds,axis,keys)
            for k in keys:
                r=hrows[k];ps=params(r)
                for g in ('main','phase'):
                    for c in r['table'][g]:
                        for m in METRICS:
                            primary.append(dict(dataset=ds,split='validation',use='正文单变量',axis='lr_multiplier' if axis=='lr_factor' else axis,anchor=anchor,run=k,
                                **{x:ps[x] for x in ('capacity','lr_multiplier','dropout','weight_decay')},section=g,condition=c,metric=m,value=r['table'][g][c][m]))
    csvwrite(OUT/'03_hyperparameters_validation.csv',primary+history,fields)
    # Main sheets strictly separate axes and anchors; historical material retained below in its own sheet.
    notes=['四数据集validation单因素；speed保留历史固定锚点，flow使用最终支持方案锚点。',
           'MAE/RMSE/MAPE越低越好；第一名加粗、第二名下划线，仅在同数据集、同锚点、同单因素曲线内排名。',
           '当前新flow实验未进行test，不填充或沿用其他配置test值。历史组合数据附于最后一张表，不作为单变量证据。']
    n=book_from_cells(OUT/'03_hyperparameters_validation.xlsx',primary,notes,['dataset','axis','section'],['run'])
    wb=load_workbook(OUT/'03_hyperparameters_validation.xlsx');ws=wb.create_sheet('历史配置与组合_长表');ws.append(fields)
    for r in history:ws.append([float(r[f]) if f=='value' else r[f] for f in fields])
    wb.save(OUT/'03_hyperparameters_validation.xlsx')
    abrows={ds+'_Full':dict(r,label='Full') for ds,r in ap['anchors'].items()};missing=[];recovery=None
    rp=ROOT/'configs/paper/ablation_recovery_20261008/plan.json'
    if rp.exists():
        recovery=read(rp);rr=ROOT/'artifacts/experiments/ablation_recovery_20261008'/recovery['plan_sha256']
    for k,e in ap['configs'].items():
        done=ar/(k+'_complete.json');plan=ap['plan_sha256']
        if not done.exists() and recovery and (rr/(k+'_complete.json')).exists():done=rr/(k+'_complete.json');plan=recovery['plan_sha256']
        if not done.exists():missing.append(dict(run=k,dataset=e['dataset'],label=e['label'],reason='OOM; isolated retry pending'));continue
        d=read(done);report=Path(d['report_path']);r=read(report)
        row=dict(dataset=e['dataset'],label=e['label'],config=e['path'],report=str(report.relative_to(ROOT)),checkpoint=str(report.with_name('best.pt').relative_to(ROOT)),table=r['development_validation_table'])
        assert r['model_method_contract']['paired_ablation']==e['ablation']
        audits['ablations'][k]=audit_row(row,e,plan);abrows[k]=row
    acells=[];deltas=[]
    for k,r in abrows.items():
        ds=r['dataset'];ref=ap['anchors'][ds]['table']
        for g in ('main','phase'):
            for c in ref[g]:
                for f,v in ref[g][c].items():
                    if f.endswith('_count'):assert r['table'][g][c][f]==v
                for m in METRICS:acells.append(dict(dataset=ds,run=k,label=r['label'],section=g,condition=c,metric=m,value=r['table'][g][c][m]))
        if r['label']!='Full':
            change={m:100*(r['table']['main']['LifeAvg'][m]/ref['main']['LifeAvg'][m]-1) for m in METRICS}
            wins=sum(ref[g][c][m]<r['table'][g][c][m] for g in ('main','phase') for c in ref[g] for m in METRICS)
            deltas.append(dict(dataset=ds,label=r['label'],lifeavg_change_percent=change,full_better_cells=wins))
    csvwrite(OUT/'05_ablations_validation.csv',acells,list(acells[0]))
    an=book_from_cells(OUT/'05_ablations_validation.xlsx',acells,
        ['四数据集固定最终配置，seed32，validation；Full对照复用。','w/o ALFRecon保留flow插值；w/o ALFTemporal–FiLM移除整条时间条件分支。',f'已完成{len(abrows)-4}/24消融；缺失项：'+str(missing),'第一加粗、第二下划线；所有可用消融均保留。'],['dataset','section'],['label'])
    wb=load_workbook(OUT/'05_ablations_validation.xlsx');ws=wb.create_sheet('LifeAvg变化百分比');ws.append(['dataset','ablation','MAE%','RMSE%','MAPE%','Full优于消融格数/33'])
    for d in deltas:ws.append([d['dataset'],d['label']]+[d['lifeavg_change_percent'][m] for m in METRICS]+[d['full_better_cells']])
    ws=wb.create_sheet('缺失状态');ws.append(['run','dataset','label','status'])
    for r in missing:ws.append(list(r.values()))
    wb.save(OUT/'05_ablations_validation.xlsx')
    write(OUT/'ablation_results_and_provenance.json',dict(rows=abrows,missing=missing,deltas=deltas,test_accessed=False))
    write(OUT/'verification.json',dict(audits=audits,new_hparam_cells=594,paper_ofat_cells=n,ablation_cells=an,missing=missing,test_accessed=False))
    lines=['# 超参数与消融结果汇总（validation，seed32）','',f'超参数14/14新训练成功，4项复用；消融{len(abrows)-4}/24成功，4项Full复用。',
        'PEMS-BAY w/o Spatial Context首次运行因GPU显存耗尽退出；失败记录保留，独立同配置重跑。空缺不是零值或模型性能。' if missing else '独立补跑已完成，消融表完整。',
        '未进行新test、未更换最终Full；新超参数结果不能替换既有配对消融的Full对照。','',
        '## 超参数：同一固定锚点下的结果','', '| 数据集/配置 | MAE↓ | RMSE↓ | MAPE↓ | 第一格/33 |','|---|---:|---:|---:|---:|']
    for k,r in hrows.items():lines.append('| '+k+' | '+' | '.join(f"{r['table']['main']['LifeAvg'][m]:.4f}" for m in METRICS)+f" | {r['first_cells']} |")
    lines+=['','PEMS04：WD1e-5和1e-3均改善三项LifeAvg且维持27格；1e-5三指标均优于1e-3。容量18偏向MAE/RMSE，但MAPE退化。',
        'PEMS08：WD1e-3三项LifeAvg改善且30→32格；LR0.5的MAPE更低但覆盖降至28格；dropout0.2改善三项并保持30格。单因素收益不能假设可叠加；本轮不启动组合或自动选型。','',
        '## 消融：相对固定Full的LifeAvg变化','', '正值表示移除组件后误差上升（支持Full）；负值表示该消融更好。','',
        '| 数据集 | 消融 | MAE变化% | RMSE变化% | MAPE变化% | Full更好格/33 |','|---|---|---:|---:|---:|---:|']
    for d in deltas:lines.append('| '+d['dataset']+' | '+d['label']+' | '+' | '.join(f"{d['lifeavg_change_percent'][m]:+.2f}" for m in METRICS)+f" | {d['full_better_cells']} |")
    lines+=['','METR-LA六项消融的三项LifeAvg均退化；PEMS-BAY首批完成五项也均退化。空间上下文补跑的最新结果及完成状态见上表，不预先假设其方向。',
        'flow结果呈现条件依赖：PEMS04/08移除Graph Refinement均改善三项LifeAvg；PEMS08移除State Age也改善三项。不能声称所有组件在所有数据集必不可少。',
        '时间条件分支对MAE/RMSE的贡献较稳定，但PEMS08移除后MAPE降低；学习式重建对PEMS08三项均重要，对PEMS04主要改善MAE/MAPE。',
        '论文应区分组件作用与配置交互，所有消融保留，单seed结果不作显著性声明。','',
        '## 文件使用','',
        '- 03_hyperparameters_validation.xlsx：四数据集单因素主表/阶段表，历史配置与组合附后；flow锚点固定最终方案，speed保留原实验锚点。',
        '- 05_ablations_validation.xlsx：四数据集全部可用消融主表/阶段表、LifeAvg变化、缺失状态。',
        '- 原01/02正式比较表及04超参数test表未改写：新消融与flow单因素尚无test，不能复制validation充当test。',
        '- verification.json与ablation_results_and_provenance.json保存审计和来源。独立重跑结束后再次运行本脚本可补齐唯一空缺。']
    (OUT/'ANALYSIS.md').write_text('\n'.join(lines)+'\n');print(json.dumps(dict(hparam_cells=n,ablation_cells=an,missing=missing),ensure_ascii=False))
if __name__=='__main__':main()
