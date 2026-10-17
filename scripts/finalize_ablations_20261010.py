"""Audit all 48 completed ablations, publish full tables and exact sign counts."""
import csv,sys
from pathlib import Path
import torch
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.run_metrla_hparams import read,write
from scripts.package_supplements_20261008 import audit_row,book_from_cells
from scripts import run_mechanism_ablations_20261008 as mechanism
from scripts import run_reconstruction_gradient_20261009 as gradient
OUT=ROOT/'paper/ablations_complete_20261010'
DS=('METRLA','PEMSBAY','PEMS04','PEMS08');MS=('MAE','RMSE','MAPE')
def main():
    torch.set_num_threads(4);OUT.mkdir(parents=True,exist_ok=True)
    original=read(ROOT/'paper/final_supplements_20261008/ablation_results_and_provenance.json');assert not original['missing']
    rows=dict(original['rows']);audits={};plans={};checks={}
    for study in (mechanism,gradient):
        p=study.check();root=study.output(p);plans[study.NAME]=p['plan_sha256'];assert read(root/'COMPLETE.json')['passed']
        assert all(j['status']=='complete' and j['exit_code']==0 for j in read(root/'queue_status.json')['jobs'])
        data=read(study.OUT/'results.json')
        for k,r in data['rows'].items():
            if r['label']=='Full':
                assert r['table']==rows[r['dataset']+'_Full']['table'];continue
            assert k not in rows;rows[k]=r;checks[k]=(p['configs'][k],p['plan_sha256'])
    assert len(rows)==52
    for k,r in rows.items():audits[k]=audit_row(r,*checks.get(k,(None,None)))
    summary={};cells=[]
    for k,r in rows.items():
        ds=r['dataset'];ref=rows[ds+'_Full']['table']
        for g in ('main','phase'):
            for c in ref[g]:
                for f,v in ref[g][c].items():
                    if f.endswith('_count'):assert r['table'][g][c][f]==v
                for m in MS:
                    value=r['table'][g][c][m];base=ref[g][c][m]
                    cells.append(dict(dataset=ds,run=k,label=r['label'],section=g,condition=c,metric=m,value=value,full_value=base,delta_percent=100*(value/base-1)))
        if r['label']=='Full':continue
        summary.setdefault(r['label'],{})[ds]={m:100*(r['table']['main']['LifeAvg'][m]/ref['main']['LifeAvg'][m]-1) for m in MS}
    assert len(summary)==12 and all(set(s)==set(DS) for s in summary.values())
    qualifying=[label for label,s in summary.items() if all(v>0 for d in s.values() for v in d.values())]
    overview=[]
    for label,s in summary.items():
        overview.append(dict(label=label,worse_metrics=sum(v>0 for d in s.values() for v in d.values()),better_metrics=sum(v<0 for d in s.values() for v in d.values()),
            equal_metrics=sum(v==0 for d in s.values() for v in d.values()),datasets_all_worse=[d for d in DS if all(v>0 for v in s[d].values())],qualifies=label in qualifying))
    wbpath=OUT/'all_ablations_validation.xlsx'
    n=book_from_cells(wbpath,cells,['12种干预×4数据集=48项完成，另有4个固定Full。seed32 validation，无缺失。',
        '判定：相同数据集固定Full，三项LifeAvg均严格更高；四数据集合计12/12才符合。未四舍五入判定。',
        '全部实验保留；单种子不作显著性声明。本文件无新test、无新训练。'],['dataset','section'],['label']);assert n==1716
    wb=load_workbook(wbpath);ws=wb.create_sheet('LifeAvg相对变化');ws.append(['消融']+[d+'/'+m+'%' for d in DS for m in MS])
    for label,s in summary.items():ws.append([label]+[s[d][m] for d in DS for m in MS])
    ws=wb.create_sheet('四数据集一致性');ws.append(['消融','变差指标数/12','改善指标数/12','相等数','三指标均变差的数据集','四数据集全部符合'])
    for r in overview:ws.append([r['label'],r['worse_metrics'],r['better_metrics'],r['equal_metrics'],','.join(r['datasets_all_worse']),r['qualifies']])
    wb.save(wbpath)
    with (OUT/'all_ablations_validation.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(cells[0]));w.writeheader();w.writerows(cells)
    write(OUT/'results_and_audit.json',dict(rows=rows,plans=plans,audits=audits,lifeavg_delta_percent=summary,overview=overview,qualifying=qualifying,qualified_count=len(qualifying),metric_cells=n,all_runs_complete=True,test_accessed=False))
    lines=['# 三轮消融完整汇总（2026-10-10）','',f'全部48项训练已完成；12种干预×4数据集，复用4项Full，共1716指标格。满足“四数据集三项LifeAvg都严格变差”的消融共 **{len(qualifying)}种**。',
        '按原始未舍入数值与各自固定Full比较；正值=消融更差，负值=消融更好。参数/数据身份、最佳epoch选择、完整计数、有限权重及零跳批审计通过；无新test/训练。','',
        '| 消融 | 变差指标/12 | 三指标均变差的数据集 |','|---|---:|---|']
    for r in overview:lines.append(f"| {r['label']} | {r['worse_metrics']} | {', '.join(r['datasets_all_worse']) or '无'} |")
    lines+=['','## 本轮两个梯度对照','', '| 数据集 | 阻断图→重建 ΔMAE/RMSE/MAPE% | 阻断时间→重建 ΔMAE/RMSE/MAPE% |','|---|---|---|']
    for d in DS:lines.append('| '+d+' | '+' | '.join('/'.join(f'{summary[l][d][m]:+.2f}' for m in MS) for l in ('Stop Graph-to-Reconstruction Gradient','Stop Temporal-to-Reconstruction Gradient'))+' |')
    lines+=['','图分支梯度对METR-LA、PEMS08有一致正向作用，但PEMS-BAY阻断后更好，PEMS04为MAE/RMSE与MAPE取舍。时间分支梯度没有在任何数据集同时改善三项LifeAvg的证据（以本轮方向比较为准）。',
        '新两项均未满足全四数据集条件。原先缺失的BAY Pointwise结果也已纳入，未改变此前结论。',
        '## 全部LifeAvg相对变化','', '| 消融 | METRLA | PEMSBAY | PEMS04 | PEMS08 |','|---|---|---|---|---|']
    for label,s in summary.items():lines.append('| '+label+' | '+' | '.join('/'.join(f'{s[d][m]:+.2f}' for m in MS) for d in DS)+' |')
    lines+=['','## 解释范围','',
        '“没有任何项在全部12指标变差”不等于完整模型没有优势，也不等于所有组件无效；它否定的是当前单seed配置下所有组件具有跨数据集、跨指标一致增益的强说法。',
        '应保留数据集依赖和指标取舍，核心结构删除、条件信息控制、梯度路径控制分别解释。不要把它们当作同等层级的独立创新。',
        '建议暂停以达到全12格变差为目标的搜索。若继续研究，优先验证稳定性或解释已发现的冲突；这不是本次自动启动的新任务。',
        '正式比较表、最终模型配置和已有test表均未修改。后续论文引用应使用本完整版本，先前中期汇总仍保留为历史快照。']
    (OUT/'ANALYSIS.md').write_text('\n'.join(lines)+'\n');print('qualifying',qualifying,'cells',n);print(overview)
if __name__=='__main__':main()
