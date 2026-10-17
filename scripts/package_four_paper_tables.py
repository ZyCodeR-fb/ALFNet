"""Four consistent paper data workbooks, from completed saved results only."""
import copy,csv,json,math,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts import test_all_final_20261001 as run
from scripts.collect_all_final_test import choose
from scripts.collect_hparams_final import params,GROUPS,METRICS
from scripts.organize_hparam_ofat import ANCHORS,AXES,expanded,ALLOWED
from scripts.run_metrla_hparams import read,write,sha
OUT=ROOT/'paper/final_four_tables_20261001'


def style(ws,metric_start):
    ws.freeze_panes=f'{get_column_letter(metric_start)}2'
    ws.sheet_view.zoomScale=85
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=16 if i>=metric_start else 22
    ws.column_dimensions['A'].width=32


def check_save(wb,path,checks):
    wb.save(path);actual=load_workbook(path)
    for sheet,i,j,value,rank in checks:
        c=actual[sheet].cell(i,j);assert abs(c.value-value)<1e-10,(sheet,i,j)
        if rank is not None:
            assert c.font.bold==(rank==1) and c.font.underline==('single' if rank==2 else None)
    return len(checks)


def workbook(notes):
    wb=Workbook();wb.active.title='说明'
    for n in notes:wb.active.append([n])
    wb.active.column_dimensions['A'].width=125
    for row in wb.active:row[0].alignment=Alignment(wrap_text=True,vertical='top')
    return wb


def comparisons(split,models_by_ds,p,selected,summary):
    filename='01_comparison_validation' if split=='validation' else '02_comparison_test'
    notes=[f'五数据集seed32对比实验：{split}；每个数据集一个ALFNet权重，validation和test用相同配置/权重。',
        '最终配置依上一轮预先写定的test探索规则选定；test参与配置选择，不可表述为未参与选择的独立确认。原validation选定版保留在../all_final_test_20261001/。',
        '第一名加粗，第二个不同数值下划线，按未舍入数值排名；MAPE单位%，MAE/RMSE为数据集原单位，不跨speed/flow平均。',
        'MoGERNN既有训练patience3；SLPF单阶段、STBP离线结构编码适配。Merlin两阶段预算和train-K教师特权保持原披露。完整审计../BASELINE_FAIRNESS_20260928.md。',
        'ST-FiT保留源Gumbel随机评估；validation采用原训练报告，test固定评估seed，历史validation不保证逐值复现。']
    wb=workbook(notes);checks=[];flat=[];md=['# '+filename,'',*notes,'']
    hp=wb.create_sheet('最终配置及权重');headers=['dataset','run',*AXES,'graph_lr','temporal_lr','gate_lr','config','checkpoint','checkpoint_sha256'];hp.append(headers)
    for ds,k in selected.items():
        r=p['entries'][k];v=dict(dataset=ds,run=k,**params(r),config=r['config'],checkpoint=r['checkpoint'],checkpoint_sha256=r['checkpoint_sha256'])
        hp.append([v[x] for x in headers])
    for ds,models in models_by_ds.items():
        summary.setdefault(ds,{})[split]=dict(first={},lifeavg={},lifeavg_ranks={})
        for g,conditions in GROUPS.items():
            ws=wb.create_sheet(ds+'_'+('主表' if g=='main' else '阶段表'));keys=[(c,m) for c in conditions for m in METRICS]
            head=['模型']+[c+'/'+m for c,m in keys];ws.append(head)
            for cell in ws[1]:cell.font=Font(bold=True)
            md += ['## '+ws.title,'','| '+' | '.join(head)+' |','|'+'|'.join(['---']+['---:']*len(keys))+'|']
            first=0
            for ri,(name,r) in enumerate(models.items(),2):
                ws.cell(ri,1,name);texts=[]
                for ci,(c,m) in enumerate(keys,2):
                    value=r['table'][g][c][m];assert math.isfinite(value)
                    rank=1+len({v['table'][g][c][m] for v in models.values() if v['table'][g][c][m]<value})
                    cell=ws.cell(ri,ci,value);cell.number_format='0.000';cell.font=Font(bold=rank==1,underline='single' if rank==2 else None)
                    if name=='ALFNet':
                        first+=rank==1
                        if c=='LifeAvg':summary[ds][split]['lifeavg'][m]=value;summary[ds][split]['lifeavg_ranks'][m]=rank
                    checks.append((ws.title,ri,ci,value,rank));flat.append([ds,name,g,c,m,value,rank])
                    t=f'{value:.3f}';texts.append(f'**{t}**' if rank==1 else f'<u>{t}</u>' if rank==2 else t)
                md.append('| '+name+' | '+' | '.join(texts)+' |')
            summary[ds][split]['first'][g]=first;style(ws,2);md.append('')
    count=check_save(wb,OUT/(filename+'.xlsx'),checks)
    (OUT/(filename+'.md')).write_text('\n'.join(md)+'\n')
    with (OUT/(filename+'.csv')).open('w',newline='') as f:w=csv.writer(f);w.writerow(['dataset','model','section','condition','metric','value','rank']);w.writerows(flat)
    return count


def hyperparameters(split,allrows,p,selected):
    filename='03_hyperparameters_validation' if split=='validation' else '04_hyperparameters_test'
    notes=[f'ALFNet四轴单变量实验：{split}完整主表/阶段表；论文正文使用各sheet上部四组，其他组合在下部灰色区，不放正文。',
        '每数据集四组共用原固定参考点，非最终test选定配置。每组只改变容量/LR倍率/dropout/WD之一，其余配置、seed32、验证LifeAvg MAE选checkpoint准则相同。',
        '参考点在四组重复列出仅方便作图，不是重复试验；第一列“正文单变量/非正文组合”明确区分，组合划分相对本表固定参考点。',
        '容量为每空间分区token容量，固定32分区；LR倍率同比例作用三个参数组；dropout为crossformer_dropout，WD为AdamW weight_decay。',
        'test曲线与全部配置已被观察并参与最终选型，是探索数据；论文参数选择主依据应如实注明validation及后续test探索过程。',
        '单变量四轴固定参考点：METRLA(8,1,.1,1e-4),BAY(16,1,.1,1e-4),03(20,1,.1,1e-4),04(10,.5,.1,1e-4),08(6,.5,.1,1e-4)。']
    wb=workbook(notes);registry=wb.create_sheet('全部67配置索引');registry.append(['dataset','run',*AXES,'改变轴','类型','最终选定','config','checkpoint'])
    checks=[];flat=[];covered=set();classification={};structure=[]
    for ds,anchor in ANCHORS.items():
        keys=[k for k,e in p['entries'].items() if e['dataset']==ds and e['model']=='ALFNet'];ps={k:params(p['entries'][k]) for k in keys};ap=ps[anchor]
        reference=expanded(p['entries'][anchor]);changes={k:[a for a in AXES if ps[k][a]!=ap[a]] for k in keys}
        for k in keys:
            q=expanded(p['entries'][k]);assert {x for x in reference.keys()|q.keys() if reference.get(x)!=q.get(x)}<=ALLOWED
            category='正文单变量' if len(changes[k])<=1 else '非正文组合';classification[k]=category
            registry.append([ds,k,*[ps[k][a] for a in AXES],'+'.join(changes[k]) or '参考点',category,k==selected[ds],p['entries'][k]['config'],p['entries'][k]['checkpoint']])
        groups=[(a,sorted([k for k in keys if all(ps[k][b]==ap[b] for b in AXES if b!=a)],key=lambda k:ps[k][a])) for a in AXES]
        groups.append(('combinations',[k for k in keys if len(changes[k])>1]))
        for g,cs in GROUPS.items():
            ws=wb.create_sheet(ds+'_'+('主表' if g=='main' else '阶段表'))
            ws.append([f'固定参考点：{anchor}；四组单变量在上、其他组合在下；{split}'])
            cols=['用途','实验组','run','轴取值',*AXES,'参考点','最终选定']+[c+'/'+m for c in cs for m in METRICS]
            for axis,ks in groups:
                label=AXES.get(axis,'其他组合（不放论文正文）');ws.append([label]);ws.cell(ws.max_row,1).font=Font(bold=True,color='FFFFFF');ws.cell(ws.max_row,1).fill=PatternFill('solid',fgColor='777777' if axis=='combinations' else '305496')
                ws.append(cols)
                for cell in ws[ws.max_row]:cell.font=Font(bold=True)
                for k in ks:
                    covered.add(k);cat='非正文组合' if axis=='combinations' else '正文单变量';v=ps[k];t=allrows[k]['table']
                    val=v[axis] if axis in AXES else '+'.join(changes[k])
                    ws.append([cat,label,k,val,*[v[a] for a in AXES],k==anchor,k==selected[ds]]+[t[g][c][m] for c in cs for m in METRICS])
                    for ci,(c,m) in enumerate(((c,m) for c in cs for m in METRICS),11):
                        cell=ws.cell(ws.max_row,ci);cell.number_format='0.000';checks.append((ws.title,ws.max_row,ci,t[g][c][m],None));flat.append([ds,split,cat,axis,anchor,k,*[v[a] for a in AXES],g,c,m,t[g][c][m]])
                    structure.append((ds,g,axis,k))
                ws.append([])
            style(ws,11);ws.freeze_panes='K4';ws.column_dimensions['C'].width=34
    assert len(covered)==67 and sum(x=='非正文组合' for x in classification.values())==20
    count=check_save(wb,OUT/(filename+'.xlsx'),checks)
    with (OUT/(filename+'.csv')).open('w',newline='') as f:w=csv.writer(f);w.writerow(['dataset','split','use','axis','anchor','run',*AXES,'section','condition','metric','value']);w.writerows(flat)
    return count,structure,classification


def main():
    p=run.check();root=run.output(p);assert read(root/'COMPLETE.json')['passed']
    q=read(root/'queue_status.json');assert len(q['jobs'])==117 and all(j['status']=='complete' and j['exit_code']==0 for j in q['jobs'])
    d=read(run.OUT/'all_results_and_provenance.json');vsource=ROOT/'paper/hparams_combinations_20260930/results.json';v=read(vsource)['rows']
    bsource=ROOT/'paper/hparams_combinations_20260930/results_and_provenance.json';b=read(bsource)['records']
    selected=d['exploratory_test_selected'];rows=d['rows'];validation={};test={};summary={};audits=[]
    for k,e in p['entries'].items():
        r=rows[k];assert r['passed'] and r['split']=='test' and r['checkpoint_sha256']==e['checkpoint_sha256']
        assert sha(ROOT/e['checkpoint'])==e['checkpoint_sha256']
        assert read(root/(k+'.json'))==r
        if 'reuse_test' not in e:
            gate=read(root/'validation'/(k+'.json'));assert gate['passed']
            assert read(root/'opened'/(k+'.json'))['plan_sha256']==p['plan_sha256']
            audits.append(dict(run=k,validation_gate=gate.get('gate','historical_agreement'),historical_agreement=gate.get('historical_agreement',True),validation_delta=gate['max_abs_difference']))
    for ds,k in selected.items():
        validation[ds]=dict(ALFNet=dict(v[k],selected_run=k),**{n:r for n,r in b[ds].items() if n!='ALFNet'})
        baselines={e['model']:rows[j] for j,e in p['entries'].items() if e['dataset']==ds and e['model']!='ALFNet'}
        candidates=[j for j,e in p['entries'].items() if e['dataset']==ds and e['model']=='ALFNet']
        assert choose(rows,candidates,baselines)[0]==k
        test[ds]=dict(ALFNet=rows[k],**baselines)
        for split,ms in [('validation',validation[ds]),('test',test[ds])]:
            for n,r in ms.items():
                for g,cs in GROUPS.items():
                    for c in cs:
                        for field,value in ms['ALFNet']['table'][g][c].items():
                            if field.endswith('_count'):assert r['table'][g][c][field]==value
    OUT.mkdir(parents=True,exist_ok=True)
    n1=comparisons('validation',validation,p,selected,summary);n2=comparisons('test',test,p,selected,summary)
    n3,s3,cl=hyperparameters('validation',v,p,selected);n4,s4,_=hyperparameters('test',rows,p,selected);assert s3==s4
    write(OUT/'selected_configuration_manifest.json',dict(selected=selected,entries={ds:p['entries'][k] for ds,k in selected.items()},selection_basis='predeclared exploratory test selector;same checkpoints in both splits',evaluation_plan_sha256=p['plan_sha256'],source_sha256={'validation_candidates':sha(vsource),'validation_baselines':sha(bsource),'tests':sha(run.OUT/'all_results_and_provenance.json')},summary=summary))
    write(OUT/'verification.json',dict(passed=True,test_records=117,reused_tests=23,new_tests=94,validation_gates=audits,comparison_validation_cells=n1,comparison_test_cells=n2,hparam_validation_cells=n3,hparam_test_cells=n4,same_selected_checkpoints_in_both_splits=True,same_hparam_layout_in_both_splits=True,unique_hparams=67,single_or_anchor_configs=47,other_combinations=20))
    lines=['# 四份论文实验数据表','',
        '117项test已全部完成并核验。按上一轮预先写定的test探索选型规则确定五个最终配置；validation和test使用相同配置/权重，不能分别挑选。因为test参与最终选择，本套表属于test探索选型后的结果，不可称为独立test确认。先前validation选定配置的test仍保留在../all_final_test_20261001/validation_selected_test.xlsx。','',
        '| 数据集 | 容量 | LR倍率 | Dropout | WD | validation第一格/33 | test主表第一/15 | test阶段第一/18 | test第一格/33 | test LifeAvg MAE / RMSE / MAPE |',
        '|---|---:|---:|---:|---:|---:|---:|---:|---:|---|']
    for ds,k in selected.items():
        x=params(p['entries'][k]);a=summary[ds]['validation'];t=summary[ds]['test'];lines.append(f"| {ds} | {x['capacity']} | {x['lr_multiplier']:g} | {x['dropout']:g} | {x['weight_decay']:g} | {sum(a['first'].values())} | {t['first']['main']} | {t['first']['phase']} | {sum(t['first'].values())} | "+' / '.join(f"{t['lifeavg'][m]:.4f}" for m in METRICS)+' |')
    lines += ['', '## 使用方式','',
        '- 01_comparison_validation.xlsx：最终配置的validation对比主表/阶段表。',
        '- 02_comparison_test.xlsx：同一最终配置的test对比主表/阶段表。',
        '- 03_hyperparameters_validation.xlsx：论文四轴单变量validation主表/阶段表；其他组合在各sheet下方灰色区，不放正文。',
        '- 04_hyperparameters_test.xlsx：同样分组/排序的test结果；其他组合同样置底。',
        '每个工作簿包含五数据集的主表、阶段表；对比表第一加粗、第二下划线。CSV提供原精度数值，组合结果仅作记录。47个唯一参考/单变量配置＋20个其他组合覆盖全部67项，参考点为作图在四组重复列出，不是重复训练。', '',
        '超参数固定参考点仍为METRLA容量8/LR1，BAY容量16/LR1，03容量20/LR1，04容量10/LR.5，08容量6/LR.5，drop.1/WD1e-4。这些不是新的最终配置；不能将已跑曲线重新标成围绕最终配置的单变量实验。完整配置与checkpoint来源见manifest和各workbook配置索引。','',
        '## 结果解读与取舍','',
        'METR-LA：容量7/LR1在test达到30/33第一，明显优于先前validation选择的25格；这也是最优配置跨split改变的实例。PEMS-BAY：LR2/WD1e-5在test保留22格第一，与测试前选择一致。',
        'PEMS03：没有候选达到三项test LifeAvg均优于最强基线，按预先规则使用LifeAvg Pareto集合后选容量16/LR1，5格第一。该配置test LifeAvg MAE/RMSE第一，MAPE第三；不能继续声称五数据集15项平均指标全优。',
        'PEMS04：低dropout/弱WD取得9格第一，比原配置8格多一格，但LifeAvg三项较原配置略差；这来自预先写定的覆盖优先规则，不是所有指标占优。原配置完整保留。',
        'PEMS08：LR2在test为10格第一，并显著优于先前13格validation优胜组合的test结果（后者test只有4格第一）。这提示验证/测试时段差异和单种子选型波动，不能将其归因于某单一模块。',
        '最终test合计76/165第一，METRLA超过80%，其余四个数据集未达80%。最终配置validation合计63/165第一；test LifeAvg14/15第一。单seed、重叠33格不能构成独立统计显著性证据。','',
        '基线预算/适配差异沿用既有披露；未修改模型、配置文件或权重，未重跑test，未修改论文正文。']
    (OUT/'README.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps(dict(cells=[n1,n2,n3,n4],summary=summary),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
