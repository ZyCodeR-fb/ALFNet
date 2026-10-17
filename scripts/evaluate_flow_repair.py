"""Audit nine flow repairs; test only the two PEMS08 validation finalists."""
import argparse
import copy
import csv
import os
from pathlib import Path
import subprocess
import sys
import time
from contextlib import ExitStack
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import run_flow_repair as train
from scripts import test_all_final_20261001 as oldtest
from scripts.run_metrla_hparams import read, write, sha, digest
from scripts.reselect_eight_baselines import audit, EXCLUDED
from scripts.collect_hparams_final import GROUPS, METRICS
from experiments.flow_repair_20261002 import training_adapter

NAME = 'flow_repair_test_20261002'
OUT = ROOT/'paper'/NAME
PLAN = ROOT/'configs/paper'/NAME/'plan.json'
SOURCE = ROOT/'paper/flow_repair_20261002/results.json'
KEYS = ('02_PEMS08_unit_focal', '08_PEMS08_unit_no_refiner')


def output(p):
    return ROOT/'artifacts/evaluation'/NAME/p['plan_sha256']


def workbook(rows, rankings, path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    wb = Workbook(); wb.active.title = '说明'
    wb.active.append(['同一数据集ALFNet变体内部第一加粗、第二下划线；第一格数单独对8个外部baseline计算。'])
    wb.active.append(['MAE/RMSE原单位，MAPE%；同seed32、相同协议。不同数据集不合并误差。'])
    checks = []; flat = []
    for ds in rankings:
        ks = [k for k, r in rows.items() if r['dataset'] == ds]
        for g, cs in GROUPS.items():
            ws = wb.create_sheet(ds+'_'+('主表' if g == 'main' else '阶段表'))
            fields = [(c,m) for c in cs for m in METRICS]
            ws.append(['run', '第一格/33(vs8)']+[f'{c}/{m}' for c,m in fields])
            for k in ks:
                ws.append([k, rankings[ds][k]['first_cells']]+[rows[k]['table'][g][c][m] for c,m in fields])
                for j,(c,m) in enumerate(fields, 3):
                    v = rows[k]['table'][g][c][m]
                    rank = sorted({rows[a]['table'][g][c][m] for a in ks}).index(v)+1
                    cell = ws.cell(ws.max_row,j);cell.font=Font(bold=rank==1,underline='single' if rank==2 else None)
                    cell.number_format='0.000';checks.append((ws.title,ws.max_row,j,v,rank))
                    flat.append([ds,k,g,c,m,v])
            ws.freeze_panes='C2';ws.column_dimensions['A'].width=37
    wb.save(path); saved=load_workbook(path)
    for sheet,i,j,v,rank in checks:
        cell=saved[sheet].cell(i,j)
        assert abs(cell.value-v)<1e-10 and cell.font.bold==(rank==1)
        assert cell.font.underline==('single' if rank==2 else None)
    with path.with_suffix('.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(['dataset','run','section','condition','metric','value']);w.writerows(flat)
    return len(checks)


def analyze():
    import torch
    p=train.check();root=train.output(p)
    assert read(root/'COMPLETE.json')['passed']
    assert all(j['status']=='complete' and j['exit_code']==0 for j in read(root/'queue_status.json')['jobs'])
    data=read(SOURCE);rows=data['rows'];evidence=[]
    for k,e in p['configs'].items():
        r=rows[k];report=read(ROOT/r['report']);cp=torch.load(ROOT/r['checkpoint'],map_location='cpu',weights_only=False)
        assert sha(ROOT/r['checkpoint'])==r['checkpoint_sha256']
        assert report['config_sha256']==e['sha256']==cp['config_sha256']
        assert report['plan_sha256']==p['plan_sha256']==cp['plan_sha256']
        assert report['development_validation_table']==r['table']
        assert cp['epoch']==report['best_epoch']==min((h['validation']['checkpoint_life_avg_mae'],h['epoch']) for h in report['history'])[1]
        assert all(h[s]['skipped']==0 for h in report['history'] for s in ('train','validation'))
        assert all(torch.isfinite(v).all() for v in cp['model_state'].values() if v.is_floating_point())
        assert cp['model_method_contract']==report['model_method_contract']
        anchor=p['anchors'][r['dataset']]['table']
        for g,cs in GROUPS.items():
            for c in cs:
                for f,v in anchor[g][c].items():
                    if f.endswith('_count'): assert r['table'][g][c][f]==v
        evidence.append(dict(run=k,epochs=len(report['history']),best_epoch=report['best_epoch'],checkpoint_sha256=r['checkpoint_sha256']))
    b=read(train.BASE_TABLE)['records'];ranking={}
    for ds in train.DATASETS:
        ks=[k for k,r in rows.items() if r['dataset']==ds]
        ranking[ds]=audit(rows,ks,{n:r for n,r in b[ds].items() if n not in EXCLUDED})
        assert all(ranking[ds][k]['first_cells']==rows[k]['first_cells'] for k in ks)
    OUT.mkdir(parents=True,exist_ok=True)
    cells=workbook(rows,ranking,OUT/'validation_results.xlsx')
    write(OUT/'validation_audit.json',dict(passed=True,training_plan=p['plan_sha256'],runs=evidence,
         ranking=ranking,verified_cells=cells,test_finalists=list(KEYS),source_sha256=sha(SOURCE),test_accessed=False))
    print('AUDIT_OK',cells)


def freeze():
    assert not PLAN.exists();train.check()
    audit_result=read(OUT/'validation_audit.json');assert audit_result['passed']
    rows=read(SOURCE)['rows'];entries={};files=[SOURCE,OUT/'validation_audit.json',OUT/'TEST_PLAN.md']
    import torch,yaml
    for k in KEYS:
        r=copy.deepcopy(rows[k]);cp=torch.load(ROOT/r['checkpoint'],map_location='cpu',weights_only=False)
        r.update(model='ALFNet',epoch=cp['epoch'],training_config_sha256=cp['config_sha256'],training_plan_sha256=cp['plan_sha256'])
        assert r['dataset']=='PEMS08' and cp['seed']==32
        assert sha(ROOT/r['config'])==r['training_config_sha256']
        entries[k]=r;cfg=yaml.safe_load((ROOT/r['config']).read_text())
        files.extend(ROOT/r[f] for f in ('config','checkpoint','report'))
        files.extend(ROOT/cfg['family']['experts'][e]['config_path'] for e in ('graph','patch'))
        files.extend(ROOT/cfg['dataset'][f] for f in ('data_path','unknown_nodes_path'))
    files.append(ROOT/'configs/paper/v4_data_manifests/PEMS08_run_manifest.json')
    files.extend([ROOT/'paper/all_final_test_20261001/all_results_and_provenance.json',ROOT/'paper/hparams_combinations_20260930/results_and_provenance.json'])
    for folder in ('model','layers','experiments','lib','data_provider','utils','scripts'):
        files.extend((ROOT/folder).rglob('*.py'))
    files.append(ROOT/'args.py')
    p=dict(kind=NAME,test_authorized=True,authorization='2026-10-02 user: analyze completed repairs and run test for worthwhile candidates',
           entries=entries,max_concurrent=2,training_adapter='experiments.flow_repair_20261002.training_adapter',
           scope='PEMS08 two validation finalists only; reuse old tests and do not train any model',
           files={str(f.relative_to(ROOT)):sha(f) for f in sorted(set(files))})
    p['plan_sha256']=digest(p);write(PLAN,p);print(p['plan_sha256'])


def check():
    p=read(PLAN);assert digest(p)==p['plan_sha256'] and p['test_authorized']
    assert set(p['entries'])==set(KEYS) and p['max_concurrent']==2
    for f,h in p['files'].items():assert sha(ROOT/f)==h,f
    return p


def worker(key):
    import torch
    from experiments import v4_runner
    from scripts.train_v4 import model_method_contract as unused
    p=check();r=p['entries'][key]
    cp=torch.load(ROOT/r['checkpoint'],map_location='cpu',weights_only=False)
    base_build=v4_runner.build_v4_session
    def verified_build(*args,**kwargs):
        from scripts.train_v4 import model_method_contract
        s=base_build(*args,**kwargs)
        assert model_method_contract(s)==cp['model_method_contract']
        return s
    with ExitStack() as stack:
        stack.enter_context(training_adapter())
        stack.enter_context(patch.object(v4_runner,'build_v4_session',verified_build))
        stack.enter_context(patch.object(oldtest,'check',check))
        stack.enter_context(patch.object(oldtest,'output',output))
        oldtest.worker(key)
    path=output(p)/(key+'.json');r=read(path)
    r['table']['note']='Two PEMS08 finalists fixed by validation before this test; historical test informed repair design.'
    r['method_contract']=cp['model_method_contract'];write(path,r)


def collect():
    p=check();rows={k:read(output(p)/(k+'.json')) for k in KEYS}
    old=read(ROOT/'paper/all_final_test_20261001/all_results_and_provenance.json')
    rows['PEMS08_lr_double']=old['rows']['PEMS08_lr_double']
    b={e['model']:old['rows'][k] for k,e in old['entries'].items() if e['dataset']=='PEMS08' and e['model'] not in EXCLUDED}
    for k,r in rows.items():
        assert r['passed']
        if k in KEYS:assert r['checkpoint_sha256']==p['entries'][k]['checkpoint_sha256']
        for baseline in b.values():
            for g,cs in GROUPS.items():
                for c in cs:
                    for f,v in r['table'][g][c].items():
                        if f.endswith('_count'):assert baseline['table'][g][c][f]==v
    ranking={'PEMS08':audit(rows,list(rows),b)}
    cells=workbook(rows,ranking,OUT/'test_results.xlsx')
    write(OUT/'test_results.json',dict(rows=rows,ranking=ranking,plan_sha256=p['plan_sha256'],verified_cells=cells))
    write(output(p)/'COMPLETE.json',dict(passed=True,new_tests=2,reused_control=1))
    print(ranking)


def queue():
    import fcntl
    p=check();root=output(p);root.mkdir(parents=True,exist_ok=True)
    lock=(root/'queue.lock').open('a');fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
    assert not (root/'queue_status.json').exists()
    active={};jobs=[]
    env=dict(os.environ,OMP_NUM_THREADS='4',MKL_NUM_THREADS='4',PYTHONUNBUFFERED='1')
    for k in KEYS:
        free=float(subprocess.check_output(['nvidia-smi','--query-gpu=memory.free','--format=csv,noheader,nounits','-i','0'],text=True).strip())
        assert free>20000,'Insufficient free GPU memory; do not affect existing jobs'
        log=(root/(k+'.log')).open('x')
        proc=subprocess.Popen([sys.executable,'-u',__file__,'worker','--key',k],cwd=ROOT,env=env,stdout=log,stderr=subprocess.STDOUT)
        active[k]=(proc,log);jobs.append(dict(key=k,status='running',pid=proc.pid))
    while active:
        for k,(proc,log) in list(active.items()):
            if proc.poll() is not None:
                job=next(j for j in jobs if j['key']==k)
                job.update(status='complete' if proc.returncode==0 and (root/(k+'.json')).exists() else 'failed',exit_code=proc.returncode)
                log.close();del active[k]
        write(root/'queue_status.json',dict(jobs=jobs,plan_sha256=p['plan_sha256'],updated=time.time()))
        if active:time.sleep(10)
    assert all(j['status']=='complete' for j in jobs),jobs
    collect()


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['analyze','freeze','check','worker','queue','collect']);parser.add_argument('--key');a=parser.parse_args()
    if a.mode=='worker':worker(a.key)
    elif a.mode=='check':print(check()['plan_sha256'])
    else:globals()[a.mode]()
