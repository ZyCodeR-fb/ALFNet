"""Validate and test the selected PEMS-BAY patch16 Focal checkpoint once."""
import argparse
import csv
import fcntl
import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
PLAN = ROOT / 'configs/paper/pemsbay_patch16_test_20260918.json'
HP_ROOT = ROOT / 'artifacts/experiments/pemsbay_hparams_20260917/1ec318e96fdff89dc664ac66868a0d5a34d2d0e0c5ef75d158381a3f4d8b1bc4'
RUN = next(HP_ROOT.glob('runs/patch16/PEMSBAY/ALFNet/*/seed32'))
CONFIG = ROOT / 'configs/paper/pemsbay_hparams_20260917/patch16.yaml'
MODELS = ('ALFNet', 'UIGNN', 'GinAR', 'MoGERNN', 'STFiT', 'SLPF', 'STGCN', 'STID', 'PatchSTG', 'STBP')
GROUPS = {'main': ('Unseen', 'Steady', 'Join', 'Fault', 'LifeAvg'),
          'phase': ('Cold', 'Warm', 'Ready', 'Onset', 'Decay', 'Outage')}
METRICS = ('MAE', 'RMSE', 'MAPE')
ATOL = 1e-5
RTOL = 1e-5


def read(path):
    return json.loads(Path(path).read_text())


def sha(path):
    with Path(path).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()


def write(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + '.tmp')
    temporary.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + '\n')
    temporary.replace(path)


def digest(value):
    return hashlib.sha256(json.dumps({k: v for k, v in value.items() if k != 'plan_sha256'},
                                    sort_keys=True).encode()).hexdigest()


def output(plan):
    return ROOT / 'artifacts/evaluation/pemsbay_patch16' / plan['plan_sha256']


def freeze():
    assert not PLAN.exists(), 'test plan already frozen'
    hp = read(ROOT / 'paper/hparams_pemsbay_20260917/results.json')
    assert hp['screening']['patch16']['passes_screen']
    report = read(RUN / 'run_report.json')
    checkpoint = RUN / 'best.pt'
    validation = RUN / 'development_validation_table.json'
    assert report['best_epoch'] == 77 and report['development_validation_table'] == read(validation)
    files = [ROOT / 'args.py', CONFIG, ROOT / 'configs/paper/pemsbay_hparams_20260917/patch16_patch.yaml',
             checkpoint, validation, RUN / 'run_report.json', ROOT / 'paper/hparams_pemsbay_20260917/results.json',
             ROOT / 'paper/validation_seed32/results_and_provenance.json',
             ROOT / 'paper/test_seed32/results_and_provenance.json',
             ROOT / 'configs/paper/v4_data_manifests/PEMSBAY_run_manifest.json',
             ROOT / 'data_v4/PEMS-BAY/raw_data.pkl', ROOT / 'data_v4/PEMS-BAY/unknown_nodes.npy']
    for folder in ('model', 'layers', 'lib', 'experiments', 'utils', 'data_provider', 'scripts'):
        files.extend((ROOT / folder).rglob('*.py'))
    files.extend((ROOT / 'configs/paper/final_seed32').glob('*.yaml'))
    plan = dict(kind='pemsbay_patch16_post_hparam_test', status='authorized',
                authorization='2026-09-18 user explicitly authorizes test and table update if a hyperparameter candidate beats Focal.',
                test_authorized=True, post_test_exploratory=True, candidate='patch16', dataset='PEMSBAY', model='ALFNet', seed=32,
                checkpoint=str(checkpoint.relative_to(ROOT)), checkpoint_sha256=sha(checkpoint),
                config=str(CONFIG.relative_to(ROOT)), config_sha256=sha(CONFIG),
                validation_table=str(validation.relative_to(ROOT)), validation_table_sha256=sha(validation),
                training_plan_sha256=report['plan_sha256'], selected_epoch=report['best_epoch'],
                selection='patch16 passed pre-specified validation screen; no test-based candidate selection',
                validation_tolerance=ATOL, validation_relative_tolerance=RTOL,
                source_validation_tables='paper/validation_seed32/results_and_provenance.json retained unchanged',
                source_test_tables='paper/test_seed32/results_and_provenance.json retained unchanged',
                files={str(f.relative_to(ROOT)): sha(f) for f in sorted(set(files))})
    plan['plan_sha256'] = digest(plan)
    write(PLAN, plan)
    print(plan['plan_sha256'])


def check():
    plan = read(PLAN)
    assert digest(plan) == plan['plan_sha256'] and plan['test_authorized']
    for name, expected in plan['files'].items():
        assert sha(ROOT / name) == expected, name
    return plan


def agree(actual, expected):
    deltas = []
    for group in ('main', 'phase'):
        assert actual[group].keys() == expected[group].keys()
        for condition, values in expected[group].items():
            for key, value in values.items():
                if key.endswith('_count'):
                    assert actual[group][condition][key] == value, (group, condition, key)
                elif key in METRICS:
                    assert math.isfinite(actual[group][condition][key])
                    delta = abs(actual[group][condition][key] - value)
                    deltas.append(delta)
                    assert delta <= ATOL + RTOL * abs(value), (group, condition, key, delta)
    return max(deltas)


def session_and_checkpoint(plan):
    import torch
    from experiments.v4_runner import build_v4_session
    checkpoint = torch.load(ROOT / plan['checkpoint'], map_location='cpu', weights_only=False)
    session = build_v4_session(ROOT / plan['config'], device='cuda:0', run_seed=32, num_workers=0)
    assert checkpoint['dataset'] == session.dataset == 'PEMSBAY'
    assert checkpoint['model'] == session.model_name == 'ALFNet'
    assert checkpoint['seed'] == session.run_seed == 32
    assert checkpoint['config_sha256'] == plan['config_sha256']
    assert checkpoint['plan_sha256'] == plan['training_plan_sha256']
    assert checkpoint['epoch'] == plan['selected_epoch']
    assert checkpoint['seed_contract'] == session.seed_contract
    assert checkpoint['scaler'] == session.scaler
    assert checkpoint['static_asset_hash'] == session.train_context.static_asset_hash()
    assert checkpoint['train_exposure_hash'] == session.train_context.exposure_hash()
    assert checkpoint['validation_exposure_hash'] == session.validation_context.exposure_hash()
    session.model.load_state_dict(checkpoint['model_state'], strict=True)
    session.model.eval()
    return session, checkpoint


def validation():
    import torch
    from experiments.v4_validation_collector import collect_validation_table, cell_count
    from lib.v4_seeds import seed_process
    plan = check()
    root = output(plan)
    path = root / 'validation.json'
    if path.exists():
        old = read(path)
        assert old['plan_sha256'] == plan['plan_sha256'] and old['passed']
        print('ALREADY COMPLETE validation')
        return
    torch.set_num_threads(4)
    started = time.time()
    session, checkpoint = session_and_checkpoint(plan)
    seed_process(session.seed_contract['derived']['training_runtime'])
    table = collect_validation_table(session, epoch=int(checkpoint['epoch']))
    delta = agree(table, read(ROOT / plan['validation_table']))
    assert cell_count(table) == 33
    write(path, dict(plan_sha256=plan['plan_sha256'], checkpoint_sha256=plan['checkpoint_sha256'],
                     passed=True, split='validation', best_epoch=checkpoint['epoch'], max_abs_difference=delta,
                     table=table, seconds=time.time()-started,
                     peak_reserved_gib=torch.cuda.max_memory_reserved()/1024**3))


def test():
    import torch
    from data_provider.v4_lifecycle_loader import V4LifecycleDataset
    from experiments.v4_runner import to_device
    from experiments.v4_validation_collector import ValidationCellCollector, cell_count
    from lib.lcbt_v4 import build_v4_teacher_forcing_mask
    from lib.runtime_context import build_runtime_context
    from lib.v4_seeds import seed_process
    from torch.utils.data import DataLoader
    from utils.model_output import unpack_model_output
    plan = check()
    root = output(plan)
    assert read(root / 'VALIDATION_COMPLETE.json')['plan_sha256'] == plan['plan_sha256']
    assert read(root / 'TEST_OPENED.json')['plan_sha256'] == plan['plan_sha256']
    path = root / 'test.json'
    if path.exists():
        old = read(path)
        assert old['plan_sha256'] == plan['plan_sha256'] and old['passed']
        print('ALREADY COMPLETE test')
        return
    torch.set_num_threads(4)
    started = time.time()
    session, checkpoint = session_and_checkpoint(plan)
    seed_process(session.seed_contract['derived']['training_runtime'])
    context = build_runtime_context('PEMSBAY', 'test', seed=32)
    assert context.static_asset_hash() == checkpoint['static_asset_hash']
    scaler = dict(checkpoint['scaler'])
    dataset = V4LifecycleDataset(context, ROOT / session.args.data_path, input_len=session.args.input_len,
                                 pred_len=session.args.pred_len, slide_step=session.args.slide_step, freq=session.args.freq,
                                 seed=session.seed_contract['derived']['train_schedule'], epoch=0, scaler=scaler,
                                 validation_seed=session.args.lcbt_validation_seed)
    loader = DataLoader(dataset, batch_size=session.train_loader.batch_size, shuffle=False, num_workers=0)
    collector = ValidationCellCollector(context, scaler, allow_phase='test')
    with torch.no_grad():
        for index, batch in enumerate(loader):
            batch = to_device(batch, session.device)
            forced = build_v4_teacher_forcing_mask(batch, session.seed_contract['derived']['teacher_forcing'],
                                                    int(checkpoint['epoch']), False)
            assert not forced.any()
            prediction, _ = unpack_model_output(session.model.forward_v4(batch, context, epoch=int(checkpoint['epoch']),
                                                                          teacher_forcing_mask=None))
            assert torch.isfinite(prediction).all()
            collector.update(batch, prediction)
            if index % 100 == 0:
                print('TEST_BATCH', index, len(loader), flush=True)
    table = collector.table()
    table.pop('development_validation_only', None)
    table.update(split='test', note='Selected patch16 after validation-only screen; authorized post-hoc confirmation test.')
    assert cell_count(table) == 33
    previous = read(ROOT / 'paper/test_seed32/results_and_provenance.json')['records']['PEMSBAY']['ALFNet']
    assert context.exposure_hash() == previous['test_exposure_hash']
    assert scaler == previous['scaler']
    for group, rows in GROUPS.items():
        for condition in rows:
            for key, value in previous['table'][group][condition].items():
                if key.endswith('_count'):
                    assert table[group][condition][key] == value
    write(path, dict(plan_sha256=plan['plan_sha256'], checkpoint_sha256=plan['checkpoint_sha256'], passed=True,
                     dataset='PEMSBAY', model='ALFNet', seed=32, split='test', best_epoch=checkpoint['epoch'],
                     scaler=scaler, teacher_forcing_positions=0, test_exposure_hash=context.exposure_hash(),
                     test_samples=len(dataset), table=table, seconds=time.time()-started,
                     peak_reserved_gib=torch.cuda.max_memory_reserved()/1024**3))


def emit_tables():
    plan = check()
    root = output(plan)
    candidate = read(root / 'test.json')
    assert candidate['passed'] and candidate['plan_sha256'] == plan['plan_sha256']
    validation_candidate = read(root / 'validation.json')
    assert validation_candidate['passed']
    old_validation = read(ROOT / 'paper/validation_seed32/results_and_provenance.json')
    old_test = read(ROOT / 'paper/test_seed32/results_and_provenance.json')
    validation_records = old_validation['records']
    validation_records = {ds: dict(records) for ds, records in validation_records.items()}
    validation_records['PEMSBAY']['Focal'] = {'table': validation_candidate['table'], 'best_epoch': plan['selected_epoch'],
                                               'checkpoint_sha256': plan['checkpoint_sha256'], 'selected_variant': 'patch16'}
    test_records = {ds: dict(records) for ds, records in old_test['records'].items()}
    test_records['PEMSBAY']['ALFNet'] = candidate
    for records, focal_key in ((validation_records, 'Focal'), (test_records, 'ALFNet')):
        for ds, models in records.items():
            reference = models[focal_key]['table']
            for record in models.values():
                for group, rows in GROUPS.items():
                    for condition in rows:
                        for key, value in reference[group][condition].items():
                            if key.endswith('_count'):
                                assert record['table'][group][condition][key] == value
    validation_out = ROOT / 'paper/validation_seed32_hparam_patch16_20260918'
    test_out = ROOT / 'paper/test_seed32_hparam_patch16_20260918'
    publish(validation_out, validation_records, 'Focal', 'validation', plan)
    publish(test_out, test_records, 'ALFNet', 'test', plan)
    write(root / 'TABLES_COMPLETE.json', dict(plan_sha256=plan['plan_sha256'], validation_dir=str(validation_out.relative_to(ROOT)),
                                              test_dir=str(test_out.relative_to(ROOT)), candidate=plan['candidate']))


def publish(out, records, focal_key, split, plan):
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    notes = [f'范围：METR-LA、PEMS-BAY，seed32，Focal＋九基线，{split}结果。',
             'PEMS-BAY Focal 行替换为容量16的patch16候选；METR-LA及九个基线复用原版同一split记录。',
             '候选由查看过原test后的超参数探索产生，先经validation筛选；本表的test为明确授权的后续确认，不是未受观察影响的独立test。',
             '每列按未舍入值排名，第一名加粗，第二个不同数值下划线；精确并列共享名次。',
             'MAE/RMSE为原速度单位，MAPE为%；不是多种子均值或独立显著性检验。']
    ws = wb.active
    ws.title = '说明'
    for note in notes:
        ws.append([note])
    ws.column_dimensions['A'].width = 130
    for row in ws:
        row[0].alignment = Alignment(wrap_text=True)
    markdown = [f'# seed32 {split}：PEMS-BAY patch16更新版', '', *notes, '']
    html_sections, cells, verification, rankings = [], [], [], {}
    for dataset, items in records.items():
        rankings[dataset] = {}
        for group, rows in GROUPS.items():
            title = dataset + ('主表' if group == 'main' else '阶段表')
            sheet = wb.create_sheet(title)
            columns = [(row, metric) for row in rows for metric in METRICS]
            sheet.append(['模型'] + [row + '/' + metric for row, metric in columns])
            markdown += ['## ' + title, '', '| 模型 | ' + ' | '.join(row + '/' + metric for row, metric in columns) + ' |',
                         '|' + '|'.join(['---'] + ['---:'] * len(columns)) + '|']
            html_rows = []
            rankings[dataset][group] = {}
            for row_index, model in enumerate(MODELS, 2):
                source = focal_key if model == 'ALFNet' else model
                label = 'ALFNet（Focal；PEMS-BAY spa=16）' if model == 'ALFNet' and dataset == 'PEMSBAY' else ('ALFNet（Focal）' if model == 'ALFNet' else ('ST-FiT' if model == 'STFiT' else ('SLPF（单阶段）' if model == 'SLPF' else model)))
                table = items[source]['table']
                values, places, html_values, md_values = [], [], [], []
                for column_index, (condition, metric) in enumerate(columns, 2):
                    value = table[group][condition][metric]
                    unique = sorted(set(r[focal_key if name == 'ALFNet' else name]['table'][group][condition][metric]
                                        for name, r in items.items() if name in MODELS))
                    place = unique.index(value) + 1
                    cell = sheet.cell(row_index, column_index, value)
                    cell.number_format = '0.000'
                    cell.font = Font(bold=place == 1, underline='single' if place == 2 else None)
                    values.append(value); places.append(place)
                    display = f'{value:.3f}'
                    html_values.append('<td>' + (f'<strong>{display}</strong>' if place == 1 else f'<u>{display}</u>' if place == 2 else display) + '</td>')
                    md_values.append(f'**{display}**' if place == 1 else f'<u>{display}</u>' if place == 2 else display)
                    cells.append([dataset, model, group, condition, metric, value, place])
                    verification.append((title, row_index, column_index, value, place))
                rankings[dataset][group][model] = {'first': places.count(1), 'second': places.count(2)}
                markdown.append('| ' + label + ' | ' + ' | '.join(md_values) + ' |')
                html_rows.append('<tr><th>' + label + '</th>' + ''.join(html_values) + '</tr>')
            markdown.append('')
            sheet.freeze_panes = 'B2'; sheet.column_dimensions['A'].width = 30
            for index in range(2, len(columns) + 2):
                sheet.column_dimensions[get_column_letter(index)].width = 15
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='183B56')
            headers = '<tr><th>模型</th>' + ''.join(f'<th>{condition}/{metric}</th>' for condition, metric in columns) + '</tr>'
            html_sections.append('<h2>' + title + '</h2><div class="scroll"><table>' + headers + ''.join(html_rows) + '</table></div>')
    workbook = out / f'{split}_seed32_tables.xlsx'
    wb.save(workbook)
    loaded = load_workbook(workbook)
    for sheet, row, column, value, rank in verification:
        cell = loaded[sheet].cell(row, column)
        assert abs(cell.value - value) < 1e-12
        assert cell.font.bold == (rank == 1) and cell.font.underline == ('single' if rank == 2 else None)
    assert len(verification) == 660
    (out / f'{split}_tables.md').write_text('\n'.join(markdown) + '\n')
    (out / f'{split}_tables.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><style>body{font-family:system-ui;margin:28px}.scroll{overflow:auto}table{border-collapse:collapse}td,th{border:1px solid #ccd5dd;padding:8px;white-space:nowrap;text-align:right}</style><h1>patch16更新版</h1>' + ''.join('<p>' + note + '</p>' for note in notes) + ''.join(html_sections) + '</html>')
    with (out / f'{split}_cells.csv').open('w', newline='') as file:
        writer = csv.writer(file); writer.writerow(['dataset', 'model', 'table', 'condition', 'metric', 'value', 'rank']); writer.writerows(cells)
    write(out / 'results_and_provenance.json', dict(plan_sha256=plan['plan_sha256'], split=split, candidate='patch16', records=records, rankings=rankings,
                                                      source_original_validation='paper/validation_seed32', source_original_test='paper/test_seed32'))
    write(out / 'verification.json', dict(cells=660, formats_verified=True, candidate='patch16', split=split,
                                           original_tables_preserved=True, plan_sha256=plan['plan_sha256']))


def queue():
    plan = check()
    root = output(plan)
    root.mkdir(parents=True, exist_ok=True)
    lock = (root / 'queue.lock').open('a')
    fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
    assert not (root / 'queue_status.json').exists(), 'existing queue: inspect, never duplicate'
    env = dict(os.environ, OMP_NUM_THREADS='4', MKL_NUM_THREADS='4', PYTHONUNBUFFERED='1', MPLCONFIGDIR='/tmp/mogernn-mpl')
    status = {'plan_sha256': plan['plan_sha256'], 'validation': 'running', 'test': 'sealed'}
    write(root / 'queue_status.json', status)
    for phase in ('validation', 'test'):
        if phase == 'test':
            assert read(root / 'validation.json')['passed']
            write(root / 'VALIDATION_COMPLETE.json', dict(plan_sha256=plan['plan_sha256'], passed=True, count=1))
            write(root / 'TEST_OPENED.json', dict(plan_sha256=plan['plan_sha256'], authorization=plan['authorization'], opened_at=time.time()))
            status.update(validation='complete', test='running'); write(root / 'queue_status.json', status)
        log = (root / (phase + '.log')).open('x')
        process = subprocess.Popen([sys.executable, '-u', str(Path(__file__).resolve()), phase], cwd=ROOT, env=env, stdout=log, stderr=subprocess.STDOUT)
        while process.poll() is None:
            status[phase + '_elapsed_seconds'] = time.time() - process.create_time if hasattr(process, 'create_time') else None
            write(root / 'queue_status.json', status)
            time.sleep(10)
        log.close()
        assert process.returncode == 0, f'{phase} failed; no automatic retry'
    write(root / 'TEST_COMPLETE.json', dict(plan_sha256=plan['plan_sha256'], count=1, completed_at=time.time()))
    emit_tables()
    status.update(test='complete', tables='complete'); write(root / 'queue_status.json', status)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode', choices=('freeze', 'check', 'validation', 'test', 'emit_tables', 'queue'))
    arguments = parser.parse_args()
    globals()[arguments.mode]()
