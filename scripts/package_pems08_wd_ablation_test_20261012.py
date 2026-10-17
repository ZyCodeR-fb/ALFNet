"""Package all four datasets' five paired ablations with coherent run identities."""
import csv,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.run_metrla_hparams import read,write
from scripts.plot_advisor_experiments_20261012 import wide_table,DS,MS
from scripts.test_pems08_wd_ablations_20261012 import OUT
KINDS=('no_age','static_router','local_context','no_temporal_film','no_refiner')
LABELS=['ALFNet','w/o State Age','w/o Adaptive Routing','w/o Spatial Context','w/o ALFTemporal–FiLM','w/o Graph Refinement']
def main():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    new=read(OUT/'results.json');nt=new['rows'];nv=read(ROOT/'paper/pems08_wd_ablations_20261012/results.json')['rows']
    ot=read(ROOT/'paper/all_ablations_test_20261010/results_and_audit.json')['rows'];ov=read(ROOT/'paper/ablations_complete_20261010/results_and_audit.json')['rows']
    rt=read(ROOT/'paper/ablation_repeat_test_20261012/results.json')['rows']['02_PEMS04_no_age']['repeat_ablation'];rv=read(ROOT/'paper/ablation_repeat_20261012/results.json')['rows']['02_PEMS04_no_age']
    data={};deltas={};metadata={}
    for split,old,n in [('validation',ov,nv),('test',ot,nt)]:
        data[split]={};deltas[split]={};metadata[split]={}
        for d in DS:
            src=n if d=='PEMS08' else old;selected={'ALFNet':src[d+'_Full']}
            for kind,label in zip(KINDS,LABELS[1:]):selected[label]=next(v for k,v in src.items() if k.endswith('_'+d+'_'+kind))
            if d=='PEMS04':selected[LABELS[1]]=rt if split=='test' else dict(table=rv['tables']['repeat_ablation'],checkpoint=rv['checkpoint'])
            data[split][d]={name:r['table'] for name,r in selected.items()};metadata[split][d]={name:{k:r[k] for k in ('checkpoint','checkpoint_sha256','config','report') if k in r} for name,r in selected.items()}
            ref=selected['ALFNet']['table'];deltas[split][d]={}
            for label in LABELS[1:]:deltas[split][d][label]={m:100*(selected[label]['table']['main']['LifeAvg'][m]/ref['main']['LifeAvg'][m]-1) for m in MS}
        wb=Workbook();wb.remove(wb.active);cells=[]
        for d in DS:
            for g in ('main','phase'):
                ws=wb.create_sheet(d+'_'+g);cols=[(c,m) for c in data[split][d]['ALFNet'][g] for m in MS];ws.append(['Variant']+[c+'/'+m for c,m in cols])
                for name,t in data[split][d].items():
                    ws.append([name]+[t[g][c][m] for c,m in cols])
                    for j,(c,m) in enumerate(cols,2):
                        for f,v in t[g][c].items():
                            if f.endswith('_count'):assert v==data[split][d]['ALFNet'][g][c][f]
                        val=t[g][c][m];ranks=sorted(set(x[g][c][m] for x in data[split][d].values()))
                        ws.cell(ws.max_row,j).font=Font(bold=val==ranks[0],underline='single' if val==ranks[1] else None);ws.cell(ws.max_row,j).number_format='0.000'
                        cells.append([d,name,g,c,m,val])
        wb.save(OUT/('four_dataset_'+split+'.xlsx'))
        with (OUT/('four_dataset_'+split+'.csv')).open('w') as f:w=csv.writer(f);w.writerow(['dataset','variant','section','condition','metric','value']);w.writerows(cells)
    values={d:{label:t['main']['LifeAvg'] for label,t in rows.items()} for d,rows in data['test'].items()}
    (OUT/'table2.md').write_text(wide_table(values,LABELS));write(OUT/'four_dataset_results.json',dict(tables=data,provenance=metadata,delta_percent=deltas))
    lines=['# 配对结果复核','','Verification Status: VERIFIED（五项validation回放）；test为本次授权的固定权重评估。','', '| 消融 | validation变差/12 | test变差/12 | PEMS08 test变差/3 |','|---|---:|---:|---:|']
    for label in LABELS[1:]:
        counts=[sum(x>0 for d in DS for x in deltas[s][d][label].values()) for s in ('validation','test')];n=sum(x>0 for x in deltas['test']['PEMS08'][label].values())
        lines.append(f'| {label} | {counts[0]} | {counts[1]} | {n} |')
    lines.extend(['','统计解释范围：seed32的一次训练/固定最佳权重，不将方向一致称作跨种子统计显著。Full的WD选择曾参考test；该测试不属于全程未观察的独立确认。PEMS04去Age是明确保留的重复记录，其他原记录均保留。','', '解释检查11/11：①按数据集和条件分解，不用总体均值掩盖反向格；②不推导逐窗口全面占优；③历史候选/消融选择范围已披露；④无新增后处理或控制变量；⑤基率问题不适用回归指标；⑥一次重跑不能证明稳定改善；⑦五项完成且全部纳入；⑧不把多候选筛选当独立显著性；⑨不改门槛或删除相反结果；⑩仅解释配对干预范围，不推出普遍因果必需性；⑪不存在据结果倒推数据分布机制的证明。'])
    (OUT/'FOUR_DATASET_ANALYSIS.md').write_text('\n'.join(lines)+'\n');print('\n'.join(lines[:10]))
if __name__=='__main__':main()
