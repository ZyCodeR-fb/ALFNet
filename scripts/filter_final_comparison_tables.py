"""User-requested comparison display subset; preserve all raw experiment records."""
import copy
import csv
import hashlib
import json
import shutil
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'paper/final_four_tables_20261001'
EXCLUDED={'STFiT','ST-FiT','Merlin'}

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def main():
    untouched={p:sha(p) for p in OUT.glob('0[34]*') if p.is_file()}
    backup=OUT/'before_model_filter_20261001';backup.mkdir(exist_ok=True)
    audit={}
    for path in sorted(OUT.glob('0[12]*.xlsx')):
        for ext in ('.xlsx','.csv','.md'):
            f=path.with_suffix(ext)
            if f.exists() and not (backup/f.name).exists():shutil.copy2(f,backup/f.name)
        wb=load_workbook(path);checks=[];flat=[];summary={}
        notes=['展示范围：ALFNet及8个比较算法；按用户要求移除ST-FiT和Merlin，原始结果保留。',
               '第一名加粗，第二个不同数值下划线；按剩余模型的未舍入数值重新排名。',
               '最终配置与权重未改变；此前test参与配置选择的记录仍有效。']
        ws=wb['说明']
        for i in range(ws.max_row,0,-1):
            text=str(ws.cell(i,1).value or '')
            if text.startswith('ST-FiT保留'):ws.delete_rows(i)
            elif 'Merlin两阶段预算' in text:
                ws.cell(i,1,text.replace('Merlin两阶段预算和train-K教师特权保持原披露。',''))
        if not any(ws.cell(i,1).value==notes[0] for i in range(1,ws.max_row+1)):
            for note in notes:ws.append([note])
        md=['# '+path.stem,'',*notes,'']
        for ws in wb:
            if not ws.title.endswith(('主表','阶段表')):continue
            original={ws.cell(i,1).value:[ws.cell(i,j).value for j in range(2,ws.max_column+1)] for i in range(2,ws.max_row+1)}
            for i in range(ws.max_row,1,-1):
                if ws.cell(i,1).value in EXCLUDED:ws.delete_rows(i)
            assert ws.max_row==10
            head=[c.value for c in ws[1]];md+=['## '+ws.title,'','| '+' | '.join(head)+' |','|'+'|'.join(['---']+['---:']*(len(head)-1))+'|']
            ds,group=ws.title.split('_');group='main' if group=='主表' else 'phase';first=0
            for i in range(2,ws.max_row+1):
                name=ws.cell(i,1).value;assert name not in EXCLUDED;shown=[]
                assert [ws.cell(i,j).value for j in range(2,ws.max_column+1)]==original[name]
                for j in range(2,ws.max_column+1):
                    cell=ws.cell(i,j);values=sorted({ws.cell(k,j).value for k in range(2,ws.max_row+1)});rank=values.index(cell.value)+1
                    font=copy.copy(cell.font);font.bold=rank==1;font.underline='single' if rank==2 else None;cell.font=font
                    checks.append((ws.title,i,j,cell.value,rank));c,m=head[j-1].split('/');flat.append([ds,name,group,c,m,cell.value,rank])
                    first+=name=='ALFNet' and rank==1
                    t=f'{cell.value:.3f}';shown.append(f'**{t}**' if rank==1 else f'<u>{t}</u>' if rank==2 else t)
                md.append('| '+name+' | '+' | '.join(shown)+' |')
            summary[ws.title]=first;md.append('')
        wb.save(path);actual=load_workbook(path)
        for sheet,i,j,v,rank in checks:
            cell=actual[sheet].cell(i,j);assert cell.value==v and cell.font.bold==(rank==1) and cell.font.underline==('single' if rank==2 else None)
        assert len(checks)==1485
        with path.with_suffix('.csv').open('w',newline='') as f:
            w=csv.writer(f);w.writerow(['dataset','model','section','condition','metric','value','rank']);w.writerows(flat)
        path.with_suffix('.md').write_text('\n'.join(md)+'\n')
        audit[path.name]=dict(verified_cells=len(checks),models_per_dataset=9,alfnet_first_cells=summary)
    assert all(sha(p)==h for p,h in untouched.items())
    (OUT/'model_filter_verification.json').write_text(json.dumps(dict(passed=True,excluded=sorted(EXCLUDED),workbooks=audit,hyperparameter_files_unchanged=True,raw_results_unchanged=True),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(audit,ensure_ascii=False))

if __name__=='__main__':main()
